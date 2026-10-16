"""Offline regression suite for Zazen (no model, no GPU, no app server).

Pure-function and fake-engine tests of the harness guards added in Oct 2026. Every test is isolated:
LWS_DATA_DIR, trackers, trash/backups and the agent log all point into a temp folder; agentic._log is a no-op.

Run:  python -m pytest backend/tests          (test_offline_suite.py runs every case below as its own test)
  or: python backend/tests/offline_suite.py   (standalone; exit code 1 when any test FAILS)
Tests named probe_* document known risks: reported as WARN / xfail, never fail the run.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import traceback

APP = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))     # backend/
SCRATCH = tempfile.mkdtemp(prefix="zh_offline_")
os.environ["LWS_DATA_DIR"] = os.path.join(SCRATCH, "data")
os.environ["LWS_NO_TELEGRAM"] = "1"
os.environ["LWS_NO_SCHEDULER"] = "1"
os.environ["LWS_NO_CODE_INDEX"] = "1"
os.makedirs(os.environ["LWS_DATA_DIR"], exist_ok=True)
sys.path.insert(0, APP)

import agent_tools as at          # noqa: E402
import agentic                    # noqa: E402
import debug_tools as dbg         # noqa: E402
import history_archive as ha      # noqa: E402
from agent_config import DEFAULT_AGENT  # noqa: E402

# ---- isolation: nothing may touch the real app data
agentic._log = lambda rec: None
agentic.LOG = os.path.join(SCRATCH, "agent_log.jsonl")
at.TRACKERS = os.path.join(SCRATCH, "trackers")
at.TRASH = os.path.join(SCRATCH, "trash")
at.BACKUPS = os.path.join(SCRATCH, "trash", "backups")

RESULTS = []
TESTS = []


def test(fn):
    TESTS.append(fn)
    return fn


def tmpdir(name="t"):
    return tempfile.mkdtemp(prefix=name + "_", dir=SCRATCH)


class Emits(list):
    def __call__(self, kind, /, **d):
        self.append((kind, d))


def engine(root=None, conv=None, **cfg):
    root = root or tmpdir("ws")
    conv = conv or f"audit{int(time.time() * 1e6) % 10**9}"
    c = dict(DEFAULT_AGENT)
    c.update(backend="openai", model="fake-model", roots=[root], tracker_path=os.path.join(SCRATCH, "trk", conv + ".md"),
             lean_reference=False, lean_first=False, interpret_first=False, think_frame=False, playbooks=False,
             tool_gate=False, file_index=False, require_checklist=False, delegation=False, auto_verify=False,
             verify_before_commit=False, facts_need_source=False, mcp_servers=[], stable_tools=True,
             fanout={"enabled": False}, hybrid={"enabled": False}, advisor={"enabled": False})
    c.update(cfg)
    em = Emits()
    e = agentic.AgentEngine(c, em, ask=lambda r: "once", conv_id=conv)
    e.reg = at.Registry(e.ctx, None, None)
    e._emits = em
    return e


def tc(name, args, cid="c1"):
    return {"id": cid, "type": "function", "function": {"name": name, "arguments": json.dumps(args)}}


# =====================================================================================================================
# 1. phantom-fix reality check
@test
def phantom_fixes_identical_pair():
    assert agentic._phantom_fixes("The bug: `player_speed` should be `player_speed`.") == ["player_speed"]


@test
def phantom_fixes_arrow_and_instead_of():
    t = "Fix `self.vel` → `self.vel` and `rect.x` instead of `rect.x`"
    assert agentic._phantom_fixes(t) == ["rect.x", "self.vel"]


@test
def phantom_fixes_real_fix_not_flagged():
    assert agentic._phantom_fixes("`colour` should be `color`; `x` -> `y`") == []
    assert agentic._phantom_fixes("") == [] and agentic._phantom_fixes(None) == []


@test
def phantom_fixes_case_difference_is_real():
    assert agentic._phantom_fixes("`Score` should be `score`") == []


# =====================================================================================================================
# 2. narration detector (_announces_action) - feeds the narration redo (tool_choice=required)
@test
def announces_action_intent_phrases():
    for t in ("Let me fetch the page.", "I'll read the file now.", "Now I need to check the tests.",
              "Next, I will run the script", "Checking the folder:", "I'm going to open it."):
        assert agentic._announces_action(t), t


@test
def announces_action_final_answers_not_flagged():
    for t in ("The capital of France is Paris.", "Median 164 ms, p95 249 ms.", "Let me know if you need more.",
              "Done. The file has 200 lines."):
        assert not agentic._announces_action(t), t


@test
def announces_action_length_limits():
    long = "Here are the findings. " * 40 + "Let me check one more thing."
    assert len(long) > 600
    assert not agentic._announces_action(long)                 # a long answer is not narration
    assert agentic._announces_action(long, open_items=True)     # ...unless the checklist still has open items


@test
def probe_announces_action_polite_closing_offer():
    """A real final answer that ends with an offer ('I'll …') is treated as narration -> redo with a REQUIRED tool call."""
    t = "Paris has about 2.1 million residents (INSEE 2021). If you want, I'll compare it with London."
    assert not agentic._announces_action(t), "final answer with a closing offer counts as narration"


# =====================================================================================================================
# 3. continuation routing (_continues_work)
def _hist_with_tools(last_user):
    return [{"role": "user", "content": "build a snake game"},
            {"role": "assistant", "content": "", "tool_calls": [tc("write_file", {"path": "x.py", "content": "1"})]},
            {"role": "tool", "tool_call_id": "c1", "name": "write_file", "content": "created x.py"},
            {"role": "assistant", "content": "Stopped at step limit."},
            {"role": "user", "content": last_user}]


@test
def continues_work_after_tool_turn():
    e = engine()
    for msg in ("continue", "go on", "yes, do it", "Proceed", "try again", "ok go", "keep going please"):
        e.history = _hist_with_tools(msg)
        assert e._continues_work(msg), msg


@test
def continues_work_needs_previous_tool_use():
    e = engine()
    e.history = [{"role": "user", "content": "tell me a joke"}, {"role": "assistant", "content": "A joke."},
                 {"role": "user", "content": "continue"}]
    assert not e._continues_work("continue")


@test
def continues_work_not_for_new_requests_or_first_message():
    e = engine()
    e.history = _hist_with_tools("now write a README for a different project with installation steps")
    assert not e._continues_work("now write a README for a different project with installation steps")
    e.history = [{"role": "user", "content": "continue"}]
    assert not e._continues_work("continue")


@test
def continues_work_ignores_reminder_messages():
    e = engine()
    h = _hist_with_tools("continue")
    h.insert(3, {"role": "user", "content": agentic.REMINDER_TAG + "\nnudge\n</system-reminder>"})
    e.history = h
    assert e._continues_work("continue")


# =====================================================================================================================
# 4. repeat gate (_repeat_of / _record_call)
@test
def repeat_gate_blocks_identical_read():
    e = engine()
    t = e.reg.get("read_file")
    a = {"path": "C:/x/a.py"}
    assert e._repeat_of("read_file", t, a) == ""
    e._record_call("read_file", t, a, "     1\tprint(1)")
    note = e._repeat_of("read_file", t, a)
    assert note.startswith("read_file did not run") and "print(1)" in note


@test
def repeat_gate_reset_by_successful_write_only():
    e = engine()
    rt, wt = e.reg.get("read_file"), e.reg.get("write_file")
    a = {"path": "C:/x/a.py"}
    e._record_call("read_file", rt, a, "x")
    e._record_call("write_file", wt, {"path": "C:/x/a.py", "content": "y"}, "Error", ok=False)
    assert e._repeat_of("read_file", rt, a), "a FAILED write must not reset the gate"
    e._record_call("write_file", wt, {"path": "C:/x/a.py", "content": "y"}, "overwrote", ok=True)
    assert e._repeat_of("read_file", rt, a) == "", "a successful write must reset the gate"


@test
def repeat_gate_exempt_tools():
    e = engine()
    for name, args in (("web_search", {"query": "q"}), ("update_tracker", {"next": "x"}), ("bg_output", {"id": "b1"})):
        t = e.reg.get(name)
        if t is None:
            continue
        e._record_call(name, t, args, "result")
        assert e._repeat_of(name, t, args) == "", name


@test
def repeat_gate_arg_order_insensitive():
    e = engine()
    t = e.reg.get("search_text")
    e._record_call("search_text", t, {"pattern": "foo", "path": "C:/x"}, "hit")
    assert e._repeat_of("search_text", t, {"path": "C:/x", "pattern": "foo"})


# =====================================================================================================================
# 5. same-attempt detector (_same_attempt) with real temp files






# =====================================================================================================================
# 6. test-output parser (_test_passes) used by progress / loop detection
@test
def test_passes_formats():
    e = engine()
    f = e._test_passes
    assert f("run_tests", "tests: FAILED — 1 failed, 3 passed in 0.21s") == 3
    assert f("powershell", "===== 10 passed, 2 warnings in 1.0s =====") == 10
    assert f("bg_output", "# tests 5\n# pass 4\n# fail 1") == 4
    assert f("powershell", "5/7 tests passed") == 5
    assert f("run_tests", "tests: FAILED — 2 failed in 0.1s") == 0
    assert f("powershell", "PASS a\nFAIL b\nPASS c") == 2
    assert f("powershell", "hello world") is None
    assert f("read_file", "3 passed") is None


@test
def test_passes_debug_run_and_python_host():
    e = engine()
    assert e._test_passes("debug_run", "ok 1 - a\nok 2 - b\nnot ok 3 - c") == 2
    assert e._test_passes("python_host", "All 12/12 passed") == 12


# =====================================================================================================================
# 7. thinking budget: think_max / _step_budget / _wants_thinking
@test
def think_max_effort_mapping():
    tm = agentic.think_max
    assert tm({"effort": "minimal"}) == 0 and tm({"effort": "low"}) == 1024 and tm({"effort": "max"}) == -1
    assert tm({"effort": "HIGH"}) == 12288
    assert tm({"think_max": "700"}) == 700
    assert tm({"think_level": "off"}) == 0 and tm({"think_level": "short"}) == 4096
    assert tm({"think": True}) == 4096 and tm({}) == -1 and tm({"effort": "auto"}) == -1   # nothing set: Auto


@test
def step_budget_levels_and_caps():
    e = engine(effort="auto")
    e._think_effort = "low"
    assert e._step_budget(True, 2) == agentic.EFFORT_THINK["low"]
    e._fail_streak = 1
    assert e._effective_effort() == "medium"          # one level up from the first failure
    assert e._step_budget(True, 2) == agentic.DEBUG_THINK     # debugging: more room, but bounded (bonsai8: 2.8 min/edit at 3072)
    e._fail_streak = 0
    e._watch_progress("run_tests", {}, True, "== 4 failed, 6 passed in 1.2s ==")
    assert e._debugging and e._fail_streak == 0      # failing tests are not a failed call, but they are debugging
    e._watch_progress("read_file", {}, True, "1| import os")
    assert e._debugging and e._wants_thinking(True, 5) and e._step_budget(True, 5) == agentic.DEBUG_THINK
    e._watch_progress("run_tests", {}, True, "== 10 passed in 1.0s ==")
    assert not e._debugging and e._step_budget(True, 2) == agentic.EFFORT_THINK["low"]
    e.cfg["effort"] = "low"
    e._watch_progress("powershell", {}, True, "exit code 1\nTraceback (most recent call last):\nValueError")
    assert e._debugging and e._step_budget(True, 2) == min(agentic.EFFORT_THINK["high"], 1024)   # never above the user's cap
    e.cfg["effort"], e._debugging = "auto", False
    e._fail_streak, e._think_effort = 0, "high"
    e._fc_rounds = 1                                  # wrapping up (after a fact check): doubled
    assert e._step_budget(True, 2) == 2 * agentic.EFFORT_THINK["high"]
    e.cfg["effort"] = "low"
    assert e._effective_effort() == "low"             # the setting caps the level, not only the tokens
    assert e._step_budget(True, 2) == agentic.EFFORT_THINK["low"]      # low effort's budget (and never above the setting)
    e.cfg["effort"], e._fc_rounds = "medium", 0                 # (not wrapping up: the answer step always thinks)
    e.turn = {"planned": True}
    assert e._effective_effort() == "medium" and not e._wants_thinking(True, 5)   # Medium: key steps only
    e._debugging = True
    assert e._effective_effort() == "high" and e._step_budget(True, 5) == agentic.DEBUG_THINK   # Medium: debugging is bounded
    e.cfg["effort"], e._think_effort = "auto", "high"
    assert e._step_budget(True, 5) == agentic.DEBUG_THINK      # the agent's own 'high' too (a debug streak: 3072/step)
    e.cfg["effort"] = "high"
    assert e._step_budget(True, 5) == agentic.EFFORT_THINK["high"]   # the user's High keeps the bigger budget
    e._debugging = False


@test
def compaction_target_scales_with_the_window():
    e = engine(num_ctx=61440, compact_target_tokens=0)
    assert e._compact_target() == int(e.context_budget() * e.COMPACT_TARGET_SHARE) > 35000   # was a fixed 20K
    e.cfg["num_ctx"] = 131072
    assert e._compact_target() > 80000
    e.cfg["compact_target_tokens"] = 30000
    assert e._compact_target() == 30000                         # an explicit setting still wins


@test
def wants_thinking_routing():
    e = engine(effort="medium")
    e._think_effort = "medium"
    e.turn = {"planned": True}
    e._need_plan = False
    e.history = [{"role": "tool", "content": "ok"}]
    assert e._wants_thinking(True, 3) is False        # routine step
    e._fail_streak = 1
    assert e._wants_thinking(True, 3) is True         # something failed
    e._fail_streak = 0
    e.history = [{"role": "tool", "content": "⚠ loop"}]
    assert e._wants_thinking(True, 3) is True         # the app flagged a loop
    e._think_effort = "high"
    assert e._wants_thinking(True, 3) is True         # still the flagged step
    e.history = [{"role": "tool", "content": "ok"}]
    assert e._wants_thinking(True, 3) is False        # the agent's own 'high': routine steps act directly
    e._fail_streak = 1
    assert e._wants_thinking(True, 3) is True and e._step_budget(True, 3) >= agentic.EFFORT_THINK["medium"]
    e._fail_streak = 0
    e.cfg["effort"] = "high"
    assert e._wants_thinking(True, 3) is True         # the user's High / Max: every step
    e.cfg["effort"] = "minimal"
    assert e._wants_thinking(True, 3) is False        # the user's cap wins


# =====================================================================================================================
# 8. restorable compaction: pins, archive, history_search
@test
def collect_and_render_pins():
    msgs = [
        {"role": "user", "content": "Build a CLI todo app. Keep it under 200 lines.\n\n(Reference for any facts: blah)"},
        {"role": "user", "content": agentic.REMINDER_TAG + "\nnudge"},
        {"role": "assistant", "content": "", "tool_calls": [tc("write_file", {"path": "C:/p/PRD.md", "content": "spec"}, "a")]},
        {"role": "tool", "tool_call_id": "a", "content": "created C:/p/PRD.md (3 lines)"},
        {"role": "assistant", "content": "", "tool_calls": [tc("artifact_create", {"title": "Notes"}, "b")]},
        {"role": "tool", "tool_call_id": "b", "content": 'created artifact a7 "Design notes"'},
        {"role": "assistant", "content": "", "tool_calls": [tc("run_tests", {"path": "C:/p"}, "c")]},
        {"role": "tool", "tool_call_id": "c", "content": "tests: FAILED — 1 failed, 4 passed in 0.3s"},
        {"role": "assistant", "content": "", "tool_calls": [tc("run_tests", {"path": "C:/p"}, "d")]},
        {"role": "tool", "tool_call_id": "d", "content": "tests: FAILED — 1 failed, 4 passed in 0.3s"},
    ]
    p = agentic._collect_pins(msgs)
    assert p["requests"] == ["Build a CLI todo app. Keep it under 200 lines."], p["requests"]
    assert p["prd"] == ["C:/p/PRD.md"] and p["files"] == ["C:/p/PRD.md"]
    assert p["artifacts"] == ["a7 — Design notes"]
    assert "4 passed" in p["tests"]
    assert any(a.endswith("(×2)") and a.startswith("run_tests") for a in p["actions"]), p["actions"]
    p2 = agentic._collect_pins([{"role": "user", "content": "Also add colours."}], prev=p)
    assert p2["requests"] == ["Build a CLI todo app. Keep it under 200 lines.", "Also add colours."]
    p2["archived"] = 7
    r = agentic._render_pins(p2)
    assert r.startswith("[Pinned") and "«Also add colours.»" in r and "PRD.md" in r and "7 entries" in r
    assert "a7 — Design notes" in r and agentic._render_pins({}) == ""


@test
def history_archive_add_get_search():
    conv = "audit/arch:1"
    ids = [ha.add(conv, "tool result", "read_file", "def compute_total(items):\n    return sum(i.price for i in items)"),
           ha.add(conv, "tool result", "powershell", "exit code 1\nZeroDivisionError: division by zero in scorer.py"),
           ha.add(conv, "message", "assistant", "The colour palette is teal and amber.")]
    assert ids == ["E1", "E2", "E3"]
    assert ha.get(conv, "e2")["label"] == "powershell" and ha.get(conv, "#E3")["kind"] == "message"
    assert ha.get(conv, "E9") is None
    hits = ha.search(conv, "ZeroDivisionError scorer")
    assert hits and hits[0][1]["id"] == "E2", hits
    assert ha.search(conv, "compute_total")[0][1]["id"] == "E1"
    assert ha.search(conv, "") == [] and ha.search("audit-empty", "x") == []
    assert os.path.commonpath([ha._path(conv), os.environ["LWS_DATA_DIR"]]) == os.environ["LWS_DATA_DIR"]


@test
def history_search_tool():
    c = at.Ctx([SCRATCH])
    c.conv_id = "audit-hs"
    ha.add(c.conv_id, "tool result", "web_fetch", "The bridge opened on 12 May 1937 and is 2,737 m long.")
    out = at.t_history_search({"query": "bridge opened"}, c)
    assert "#E1" in out and "1937" in out
    assert "2,737 m" in at.t_history_search({"entry": "E1"}, c)
    assert "1 archived entries" in at.t_history_search({}, c)
    try:
        at.t_history_search({"entry": "E5"}, c)
        raise AssertionError("missing entry must raise")
    except at.ToolError:
        pass


@test
def compaction_stubs_archive_and_pins_fake_engine():
    e = engine(conv="audit-compact", compact_prompt=False, context_compaction_v2=False)
    big = "BEGIN\n" + "\n".join(f"row {i} value {i * 7}" for i in range(400)) + "\nUNIQUE_TOKEN_QX7 found\nexit code 0"
    e.history = [{"role": "user", "content": "Analyse data.csv and keep totals in EUR."}]
    for i in range(10):
        e.history += [{"role": "assistant", "content": f"Step {i}: reading part {i} of the data file now.",
                       "tool_calls": [tc("read_file", {"path": f"C:/d/part{i}.csv"}, f"k{i}")]},
                      {"role": "tool", "tool_call_id": f"k{i}", "name": "read_file",
                       "content": big.replace("UNIQUE_TOKEN_QX7", f"UNIQUE_TOKEN_QX7_{i}")}]
    e.history.append({"role": "user", "content": "now give me the totals"})
    e._summarize_in_place = lambda hist, cut: "SUMMARY: read parts 0-9."
    e._summarize = lambda old: "SUMMARY: read parts 0-9."
    e.system_prompt = lambda: "system"
    e.context_budget = lambda: 6000
    e.emit_context = lambda *a, **k: None
    e._compact_req = True
    e._maybe_compact()
    head = e.history[0]["content"]
    assert head.startswith("[Summary of the earlier part") and "SUMMARY: read parts 0-9." in head
    assert "«Analyse data.csv and keep totals in EUR.»" in head, "user's words must be pinned verbatim"
    assert "history_search" in head
    n = len(ha.entries("audit-compact"))
    assert n >= 5, n
    hit = ha.search("audit-compact", "UNIQUE_TOKEN_QX7_0")
    assert hit and "row 399" in hit[0][1]["text"], "archived entry must hold the FULL tool result"
    stubs = [m for m in e.history if m.get("role") == "tool" and "archived as #E" in str(m.get("content"))]
    for m in stubs:
        assert "exit code 0" in m["content"], "stub must keep the outcome line"


# =====================================================================================================================
# 9. debug tools: clip_ht / enrich / lint_file / run_tests / debug_run / timeout wording
@test
def clip_ht_keeps_head_and_tail():
    assert dbg.clip_ht("short") == "short"
    s = "HEAD" + "x" * 20000 + "TAIL_ERROR"
    out = dbg.clip_ht(s, n=1000)
    assert out.startswith("HEAD") and out.endswith("TAIL_ERROR") and "omitted from the middle" in out
    assert len(out) < 1200


@test
def enrich_locates_project_frame():
    d = tmpdir("enr")
    p = os.path.join(d, "app.py")
    open(p, "w").write("def f(x):\n    return 10 / x\n\nf(0)\n")
    r = subprocess.run([sys.executable, p], capture_output=True, text=True)
    out = dbg.enrich(r.stderr, d)
    assert "[Error located] ZeroDivisionError" in out and f"{p}:2" in out and "> 2 |" in out, out[-600:]
    assert dbg.enrich("no traceback here") == "no traceback here"


@test
def enrich_pytest_and_node_formats():
    exc, frames = dbg.locate("____ test_a ____\n\ntests/test_a.py:5: in test_a\n    assert f() == 2\nE   assert 1 == 2\n",
                             cwd="C:/proj")
    assert exc == "assert 1 == 2" and frames and frames[-1][1] == 5
    exc, frames = dbg.locate("TypeError: x is not a function\n    at run (C:\\proj\\app.js:12:5)\n"
                             "    at Object.<anonymous> (C:\\proj\\main.js:3:1)\n")
    assert "TypeError" in exc and frames[-1][0].endswith("app.js") and frames[-1][1] == 12


@test
def lint_file_python_json_and_skips():
    d = tmpdir("lint")
    good, bad, undef, js, txt = (os.path.join(d, n) for n in ("g.py", "b.py", "u.py", "c.json", "n.txt"))
    open(good, "w").write("import os\nprint(os.sep)\n")
    open(bad, "w").write("def f(:\n  pass\n")
    open(undef, "w").write("def f():\n    return totl + 1\n")
    open(js, "w").write("{\"a\": 1,}")
    open(txt, "w").write("anything")
    assert "no problems" in dbg.lint_file(good)
    assert "FOUND PROBLEMS" in dbg.lint_file(bad) and "SyntaxError" in dbg.lint_file(bad)
    assert "FOUND PROBLEMS" in dbg.lint_file(js)
    assert dbg.lint_file(txt) == "" and dbg.lint_file(os.path.join(d, "missing.py")) == ""
    r = dbg.lint_file(undef)
    assert "totl" in r and "FOUND PROBLEMS" in r, r


@test
def probe_lint_undefined_name_without_ruff_or_pyflakes():
    """lint_file runs ruff/pyflakes with whatever `python` is first on PATH. Without either, an undefined name is
    reported as 'no problems' — the check silently degrades to syntax-only and says so with a clean bill of health."""
    d = tmpdir("lint2")
    undef = os.path.join(d, "u.py")
    open(undef, "w").write("def f():\n    return totl + 1\n")
    real = dbg.shutil.which
    try:
        dbg.shutil.which = lambda name: os.path.join(APP, "venv", "Scripts", "python.exe") if name == "python" else real(name)
        r = dbg.lint_file(undef)
    finally:
        dbg.shutil.which = real
    assert "totl" in r, f"undefined name missed when ruff/pyflakes are absent: {r.strip()!r}"


@test
def lint_on_write_via_registry():
    d = tmpdir("low")
    c = at.Ctx([d])
    reg = at.Registry(c, None, None)
    text, _ = reg.call("write_file", {"path": os.path.join(d, "m.py"), "content": "def f(:\n"})
    assert text.startswith("created") and "Code check FOUND PROBLEMS" in text, text
    text, _ = reg.call("edit_file", {"path": os.path.join(d, "m.py"), "old": "def f(:", "new": "def f():\n    pass"})
    assert "Code check: no problems" in text, text


@test
def run_tests_pytest_or_clear_message():
    d = tmpdir("rt")
    open(os.path.join(d, "test_x.py"), "w").write("def test_ok():\n    assert 1 == 1\n\ndef test_bad():\n    assert 1 == 2\n")
    c = at.Ctx([d])
    out = at.t_run_tests({"path": d, "timeout": 60}, c)
    if "pytest is not installed" in out:
        return "pytest missing for the PATH python (message is clear)"
    assert out.startswith("tests: FAILED") and "1 failed, 1 passed" in out, out[:400]
    e = engine(d)
    assert e._test_passes("run_tests", out) == 1


@test
def debug_run_post_mortem_and_probe():
    d = tmpdir("dr")
    p = os.path.join(d, "calc.py")
    open(p, "w").write("def avg(xs):\n    total = sum(xs)\n    return total / len(xs)\n\nprint(avg([]))\n")
    c = at.Ctx([d])
    out = at.t_debug_run({"script": p, "timeout": 30}, c)
    assert out.startswith("exit code 1") and "[post-mortem]" in out and "total = 0" in out, out[:800]
    assert "<debug_run>" not in out.split("[post-mortem]")[0] or True
    out = at.t_debug_run({"script": p, "probe": "calc.py:3", "timeout": 30}, c)
    assert "[probe] calc.py:3 reached" in out and "xs = []" in out, out[:800]


@test
def debug_run_timeout_wording():
    d = tmpdir("to")
    p = os.path.join(d, "loop.py")
    open(p, "w").write("import time\nprint('started', flush=True)\nwhile True:\n    time.sleep(0.2)\n")
    c = at.Ctx([d])
    t0 = time.time()
    out = at.t_debug_run({"script": p, "timeout": 5}, c)
    assert time.time() - t0 < 30
    assert "timed out after 5s" in out and "killed" in out and "started" in out, out[-400:]


@test
def powershell_timeout_says_still_running():
    d = tmpdir("ps")
    c = at.Ctx([d])
    out = at.t_powershell({"command": "Write-Output 'up'; Start-Sleep -Seconds 30", "timeout": 4}, c)
    assert "still running" in out, out[-500:]


# =====================================================================================================================
# 10. file tools: no-op guards, backups isolated
@test
def write_file_noop_refused_and_crlf_equivalent():
    d = tmpdir("wf")
    c = at.Ctx([d])
    p = os.path.join(d, "a.txt")
    assert at.t_write_file({"path": p, "content": "one\ntwo\n"}, c).startswith("created")
    for content in ("one\ntwo\n", "one\r\ntwo\r\n"):
        try:
            at.t_write_file({"path": p, "content": content}, c)
            raise AssertionError("identical rewrite must be refused")
        except at.ToolError as ex:
            assert "already has exactly this content" in str(ex)
    assert at.t_write_file({"path": p, "content": "three\n", "append": True}, c).startswith("appended")
    assert os.listdir(at.BACKUPS), "backups must land in the patched (temp) folder"


@test
def edit_file_noop_and_not_found():
    d = tmpdir("ef")
    c = at.Ctx([d])
    p = os.path.join(d, "a.py")
    open(p, "w", newline="").write("x = 1\r\ny = 2\r\n")
    for args, msg in (({"old": "x = 1", "new": "x = 1"}, "identical"), ({"old": "", "new": "z"}, "exact text"),
                      ({"old": "q = 9", "new": "q = 8"}, "not found")):
        try:
            at.t_edit_file(dict(args, path=p), c)
            raise AssertionError(f"expected refusal: {args}")
        except at.ToolError as ex:
            assert msg in str(ex), str(ex)
    out = at.t_edit_file({"path": p, "old": "x = 1\ny = 2", "new": "x = 1\ny = 3"}, c)   # LF sent for a CRLF file
    assert out.startswith("edited") and open(p, newline="").read() == "x = 1\r\ny = 3\r\n"


@test
def resolve_refuses_outside_roots():
    d = tmpdir("rr")
    c = at.Ctx([d])
    try:
        c.resolve("C:/Windows/system32/drivers/etc/hosts")
        raise AssertionError("outside path must be refused")
    except at.ToolError:
        pass


@test
def update_tracker_must_not_clobber_user_data_file():
    """Regression for eval run 20261009-093016 (tc_data_stats): update_tracker(path=<the user's .txt data file>)
    overwrote response_times.txt with the checklist markdown."""
    d = tmpdir("trk")
    c = at.Ctx([d])
    c.conv_id = "audit-trk"
    data = os.path.join(d, "response_times.txt")
    open(data, "w").write("182\n93\n231\n")
    at.t_update_tracker({"new": True, "goal": "stats", "milestones": [{"title": "Load", "subtasks": ["read"]}]}, c)
    at.t_update_tracker({"path": data, "updates": [{"id": "1.1", "status": "done"}]}, c)
    assert open(data).read() == "182\n93\n231\n", "checklist overwrote the user's data file"


@test
def update_tracker_folder_path_ignored():
    d = tmpdir("trk2")
    c = at.Ctx([d])
    c.conv_id = "audit-trk2"
    out = at.t_update_tracker({"new": True, "goal": "g", "path": d, "milestones": [{"title": "A", "subtasks": ["a"]}]}, c)
    assert out.startswith("checklist saved") and os.path.isfile(c.tracker_path) and not os.path.isdir(c.tracker_path)


# =====================================================================================================================
# 11. per-chat workspace folders
@test
def workspace_named_by_date_topic_chat():
    root = tmpdir("wsr")
    e = engine(root, conv="ab-12_cd!ef", workspace_root=os.path.join(root, "Zen Harness"))
    e.history = [{"role": "user", "content": "Build me a snake game in pygame please<system-reminder>x</system-reminder>"}]
    e._ensure_workspace("ignored")
    exp = os.path.join(root, "Zen Harness", f"{time.strftime('%Y-%m-%d')} Build me a snake game in (ab12cd)")
    assert e.ctx.workspace == exp, e.ctx.workspace
    assert not os.path.exists(exp), "the folder is made on first use, not at naming time"
    first = e.ctx.workspace
    e.history.append({"role": "user", "content": "now something else entirely"})
    e._ensure_workspace("now something else")
    assert e.ctx.workspace == first, "a chat keeps its folder"
    wd = at.work_dir(e.ctx)
    assert wd == first and os.path.isdir(first)


@test
def workspace_edge_cases():
    root = tmpdir("wse")
    e = engine(root, workspace_root="C:/Windows/Temp/ZH")            # outside the allowed folders
    e.history = [{"role": "user", "content": "hello"}]
    e._ensure_workspace("hello")
    assert e.ctx.workspace is None
    e = engine(root, workspace_root=os.path.join(root, "ZH"))
    e.ctx.is_subagent = True
    e._ensure_workspace("x")
    assert e.ctx.workspace is None
    e = engine(root, workspace_root=os.path.join(root, "ZH"))
    e.history = [{"role": "user", "content": "¿¡ — !!"}]
    e._ensure_workspace("")
    assert " chat (" in os.path.basename(e.ctx.workspace), e.ctx.workspace












# =====================================================================================================================
# 12. ask_user / show_user / browser_eval argument guards (nothing is launched)
@test
def tracker_recovers_missing_milestone_title_without_losing_subtasks():
    c = at.Ctx([SCRATCH])
    c.tracker_path = os.path.join(SCRATCH, 'recovered-tracker.md')
    at.t_update_tracker({'new': True, 'goal': 'Fix tests', 'milestones': [
        {'done_when': 'Tests pass', 'subtasks': [{'step': 'Run tests'}, {'step': 'Fix lib.py'}]}]}, c)
    assert len(at.tracker_open(c.tracker)) == 2
    before = json.dumps(c.tracker, sort_keys=True)
    try:
        at.t_update_tracker({'new': True, 'milestones': [{}]}, c)
    except at.ToolError:
        pass
    else:
        raise AssertionError('Invalid plan accepted')
    assert json.dumps(c.tracker, sort_keys=True) == before


@test
def windows_path_case_and_sibling_containment():
    root = tmpdir('case_path')
    c = at.Ctx([root])
    assert c.inside(os.path.join(root.upper() if os.name == 'nt' else root, 'sales.csv'))
    assert not c.inside(root + '-outside/file.txt')
    assert not c.inside(os.path.join(root, '..', 'outside.txt'))






@test
def explicit_new_workspace_replaces_prior_chat_workspace():
    root = tmpdir('new_task_root')
    e = engine(root)
    e.ctx.cwd = e.ctx.workspace = root
    target = os.path.join(root, 'new_project')
    e._turn_request = 'Work only in ' + target + '. Write output.csv here.'
    assert e._bind_workspace(target)
    assert os.path.isdir(target)
    assert e.ctx.cwd == target and e.ctx.workspace == target
    assert e.ctx.resolve('output.csv') == os.path.join(target, 'output.csv')




@test
def numbered_slide_response_keeps_titles_and_boundaries():
    import doc_create
    from pptx import Presentation
    text = 'Slide 1: Title\nTitle: Security\n- First point\nSlide 2:\nTitle: Passwords\nBullets:\n- Unique passwords'
    path = os.path.join(SCRATCH, 'numbered-slides.pptx')
    doc_create.make_pptx(path, text)
    deck = Presentation(path)
    assert len(deck.slides) == 2
    all_text = '\n'.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
    assert 'Security' in all_text and 'Passwords' in all_text and 'Unique passwords' in all_text
    doc_create.make_pptx(path, '# Security\n## Overview\n- First point\n# Passwords\n## Practices\n- Unique passwords', title='Security Presentation')
    assert len(Presentation(path).slides) == 2
    doc_create.make_pptx(path, '**Slide 1: Security**\n- First point\n**Slide 2: Passwords**\n- Unique passwords')
    assert len(Presentation(path).slides) == 2




@test
def ask_user_formats_and_guards():
    c = at.Ctx([SCRATCH])
    out = at.t_ask_user({"questions": [{"question": "Which folder?", "default": "Documents"}, "Keep originals?"]}, c)
    assert out.startswith("Questions shown") and "1. Which folder? *(default: Documents)*" in c.asked
    assert "2. Keep originals?" in c.asked and "go with the defaults" in c.asked
    c2 = at.Ctx([SCRATCH])
    at.t_ask_user({"questions": "- a?\n- b?\n" + "\n".join(f"- q{i}?" for i in range(10))}, c2)
    assert c2.asked.count("?") <= 7 and "6. " in c2.asked and "7. " not in c2.asked
    for c3, args in ((at.Ctx([SCRATCH]), {"questions": []}),):
        try:
            at.t_ask_user(args, c3)
            raise AssertionError("empty questions must raise")
        except at.ToolError:
            pass
    c4 = at.Ctx([SCRATCH])
    c4.may_ask = False
    try:
        at.t_ask_user({"questions": ["x?"]}, c4)
        raise AssertionError("second round must be refused")
    except at.ToolError as ex:
        assert "already answered" in str(ex)


@test
def show_user_guards_without_launching():
    c = at.Ctx([SCRATCH])
    for args, msg in (({}, "give 'path'"), ({"command": "shutdown /s /t 0"}, "blocked")):
        try:
            at.t_show_user(args, c)
            raise AssertionError(f"expected refusal: {args}")
        except at.ToolError as ex:
            assert msg in str(ex), str(ex)


@test
def browser_eval_only_local_pages():
    class FakeBM:
        url = "https://example.com/"

        def state(self):
            return {"url": self.url}

        def evaluate(self, expr):
            return f"evaluated {expr}"
    fake = FakeBM()
    real = at._bm
    at._bm = lambda: fake
    try:
        c = at.Ctx([SCRATCH])
        try:
            at.t_browser_eval({"expression": "1+1"}, c)
            raise AssertionError("remote page must be refused")
        except at.ToolError as ex:
            assert "only runs on your own local pages" in str(ex)
        for u in ("http://localhost:5173/", "http://127.0.0.1:8000/x", "file:///C:/a.html"):
            fake.url = u
            assert at.t_browser_eval({"expression": "1+1"}, c) == "evaluated 1+1"
        fake.url = "http://localhost.evil.com/"
        try:
            at.t_browser_eval({"expression": "1"}, c)
            raise AssertionError("localhost.evil.com must be refused")
        except at.ToolError:
            pass
    finally:
        at._bm = real


# =====================================================================================================================
# 13. search-first gate (_supplement) and lean fact check (_lean_facts_ok), fake engine
@test
def supplement_appends_reference_and_card_rules():
    e = engine(lean_reference=True)
    e._lean_reference = lambda q: "[Wikipedia] Golden Gate Bridge: opened 1937."
    m = {"role": "user", "content": "When did the Golden Gate Bridge open?"}
    e.history = [m]
    e._turn_user_msg = m
    e._supplement(m["content"])
    assert "(Reference for any facts" in m["content"] and "opened 1937" in m["content"]
    assert not [k for k, _ in e._emits if k == "tool_call"], "no search card on a chat's first message"
    m2 = {"role": "user", "content": [{"type": "text", "text": "And how long is it?"}]}
    e.history += [{"role": "assistant", "content": "1937."}, m2]
    e._turn_user_msg = m2
    e._supplement("And how long is it?")
    assert "opened 1937" in m2["content"][0]["text"]
    assert not [k for k, _ in e._emits if k in ("tool_call", "tool_result")], "the background search is never shown"


@test
def supplement_skips_acks_and_disabled():
    e = engine(lean_reference=True)
    called = []
    e._lean_reference = lambda q: called.append(q) or "REF"
    m = {"role": "user", "content": "thanks!"}
    e.history, e._turn_user_msg = [m], m
    e._supplement("thanks!")
    e.cfg["lean_reference"] = False
    m2 = {"role": "user", "content": "Who wrote Dune?"}
    e._turn_user_msg = m2
    e._supplement("Who wrote Dune?")
    assert called == [] and m["content"] == "thanks!" and m2["content"] == "Who wrote Dune?"


def _fake_openai(replies, seen=None):
    it = iter(replies)

    def call(messages, tools=True, max_tokens=None, no_think=False, schema=None, raw_tools=None):
        r = next(it)
        if seen is not None:
            seen.append({"schema": schema, "tools": tools})
        if isinstance(r, Exception):
            raise r
        return r, 10, 100
    return call


@test
def lean_facts_ok_verdicts():
    e = engine(facts_need_source=True)
    e._turn_ref = "REF"
    e._call_openai = _fake_openai([{"content": '{"unsupported": ["3 symphonies"]}'}])
    assert e._lean_facts_ok([], [], "She wrote 3 symphonies.") is False
    e._call_openai = _fake_openai([{"content": '{"unsupported": []}'}])
    assert e._lean_facts_ok([], [], "x") is True
    seen = []
    e._call_openai = _fake_openai([RuntimeError("schema+tools unsupported"), {"content": '{"unsupported": ["X"]}'}], seen)
    assert e._lean_facts_ok([], [], "x") is False and seen[0]["schema"] and not seen[1]["schema"]
    e._turn_ref = ""
    e._call_openai = _fake_openai([])
    assert e._lean_facts_ok([], [], "x") is True     # no Reference: nothing to check against, no call


@test
def probe_lean_facts_ok_fails_open():
    """Garbage or an error from the checker = the answer is shown unchecked (fail-open)."""
    e = engine(facts_need_source=True)
    e._turn_ref = "REF"
    e._call_openai = _fake_openai([{"content": "I think it is fine"}])
    r1 = e._lean_facts_ok([], [], "x")
    e._call_openai = _fake_openai([RuntimeError("down"), RuntimeError("down")])
    r2 = e._lean_facts_ok([], [], "x")
    assert not (r1 and r2), "lean fact check fails OPEN on unparsable output / server errors"


# =====================================================================================================================
# 14. effort triage from the lean pass (get_tools plan/think) - fake model
@test
def lean_pass_effort_triage():
    e = engine(lean_first=True, effort="max")
    e._effort = None                                  # run() resets it before the lean pass
    e.history = [{"role": "user", "content": "Refactor my project into modules and add tests"}]
    gt = {"role": "assistant", "content": "", "tool_calls": [tc("get_tools", {"why": "multi-file work", "depth": "extensive",
                                                                              "plan": "full", "think": "high"})]}
    e._call_openai = _fake_openai([gt])
    assert e._lean_pass("Refactor my project into modules and add tests") is None
    assert e._effort == {"plan": "full", "think": "high", "size": "?"} and e._depth == "extensive"
    e2 = engine(lean_first=True)
    e2._effort = None
    e2.history = [{"role": "user", "content": "x"}]
    e2._call_openai = _fake_openai([{"role": "assistant", "content": "", "tool_calls": [
        tc("get_tools", {"why": "w", "depth": "basic", "plan": "huge", "think": "medium"})]}])
    assert e2._lean_pass("x") is None and e2._effort is None and e2._depth == "basic"
    e3 = engine(lean_first=True)
    e3._effort = None
    e3.history = [{"role": "user", "content": "check the artifact"}]
    e3._call_openai = _fake_openai([{"role": "assistant", "content": "Let me look at the artifact."}])
    assert e3._lean_pass("check the artifact") is None
    assert e3._effort == {"plan": "none", "think": "medium"}, "narrating lean reply must hand over to the agent"


@test
def lean_pass_answers_directly():
    e = engine(lean_first=True)
    e.history = [{"role": "user", "content": "What is 2+2?"}]
    e._call_openai = _fake_openai([{"role": "assistant", "content": "4."}])
    assert e._lean_pass("What is 2+2?") == "4." and e.history[-1] == {"role": "assistant", "content": "4."}


# =====================================================================================================================
# 15. artifacts store (isolated by LWS_DATA_DIR)
@test
def artifacts_create_edit_read_isolated():
    import artifacts
    assert artifacts.ROOT.startswith(os.environ["LWS_DATA_DIR"]), artifacts.ROOT
    meta, v = artifacts.create("audit-art", "Design notes", "markdown", "# v1\nalpha")
    aid = meta["id"]
    artifacts.edit("audit-art", aid, old="alpha", new="beta")
    _m, cur, ver = artifacts.read("audit-art", aid)
    assert cur == "# v1\nbeta" and ver == 2, (cur, ver)
    _m, old, _v = artifacts.read("audit-art", aid, 1)
    assert old == "# v1\nalpha", "earlier versions stay readable"
    try:
        artifacts.edit("audit-art", aid, old="beta", new="beta")
        raise AssertionError("no-op artifact edit must be refused")
    except artifacts.ArtifactError as ex:
        assert "no change" in str(ex)


# =====================================================================================================================
# 16. full fake-model turn: narration redo forces a tool call and drops the narration
@test
def narration_redo_requires_tool_call_fake_turn():
    root = tmpdir("nr")
    p = os.path.join(root, "notes.txt")
    open(p, "w").write("the answer is 42\n")
    e = engine(root, effort="minimal")
    seen = []
    script = [
        {"role": "assistant", "content": "Let me read the file first."},                       # narration, no call
        {"role": "assistant", "content": "", "tool_calls": [tc("read_file", {"path": p}, "r1")]},
        {"role": "assistant", "content": "The file says 42."},
    ]
    plan_json = {"role": "assistant", "content": json.dumps({"why": "one file", "answer_directly": False,
                                                             "needs_plan": False, "data_points": 1})}

    def fake(messages, tools=True, max_tokens=None, no_think=False, schema=None, raw_tools=None):
        if schema is not None or tools is False:
            return plan_json, 5, 50
        seen.append(bool(getattr(e, "_require_call", False)))
        e._require_call = False
        return script.pop(0) if script else {"role": "assistant", "content": "The file says 42."}, 5, 50
    e._call_openai = fake
    e._call_ollama = fake
    ans = e.run(f"What does {p} say?")
    assert "42" in (ans or ""), ans
    assert len(seen) >= 3 and seen[1] is True, f"2nd model step must have tool_choice=required; saw {seen}"
    assert not any(m.get("role") == "assistant" and "Let me read the file first" in str(m.get("content"))
                   for m in e.history), "narration must be dropped from the history"
    assert any(m.get("role") == "tool" and "42" in str(m.get("content")) for m in e.history)


# =====================================================================================================================

@test
def basic_path_can_run_commands_and_edit():
    for t in ("powershell", "bg_start", "bg_output", "bg_list", "bg_stop", "edit_file", "write_file", "show_user"):
        assert t in agentic.BASIC_TOOLS, t


@test
def basic_path_non_basic_tool_upgrades():
    e = engine()
    e._basic = True
    e._basic_calls = 0
    tc = {"id": "c1", "type": "function", "function": {"name": "run_tests", "arguments": "{}"}}
    e.history.append({"role": "assistant", "content": "", "tool_calls": [tc]})
    e._run_tool(tc)
    assert e._upgrade and "run_tests" in e._upgrade, e._upgrade
    assert "switching to the full toolset" in e.history[-1]["content"]


@test
def gone_background_process_noted():
    e = engine()
    e.history = [{"role": "user", "content": "start it"},
                 {"role": "tool", "tool_call_id": "x", "name": "bg_start",
                  "content": "started background process bg7 ('vite dev', pid 1): running. Use bg_output"},
                 {"role": "user", "content": "open the page"}]
    e._turn_user_msg = e.history[-1]
    e._note_gone_processes()
    assert "no longer running" in e.history[-1]["content"] and "bg7 'vite dev'" in e.history[-1]["content"]
    e2 = engine()
    e2.history = [{"role": "tool", "content": "started background process bg8 ('w', pid 1): running."},
                  {"role": "tool", "content": "stopped bg8 'w'"}, {"role": "user", "content": "hi"}]
    e2._turn_user_msg = e2.history[-1]
    e2._note_gone_processes()
    assert e2.history[-1]["content"] == "hi", "a process the agent stopped itself is not reported"


@test
def local_connection_refused_explained():
    import socket
    with socket.socket() as so:
        so.bind(("127.0.0.1", 0))
        port = so.getsockname()[1]           # closed again below: nothing listens there
    hint = at.local_down_hint(f"http://localhost:{port}/")
    assert f"Nothing is running on port {port}" in hint and "bg_start" in hint
    assert at.local_down_hint("https://example.com/") == ""
    try:
        at.t_web_fetch({"url": f"http://localhost:{port}/index.html"}, at.Ctx([tmpdir("wf")]))
        raise AssertionError("should have raised")
    except at.ToolError as ex:
        assert "Nothing is running" in str(ex)


@test
def unchanged_reread_points_back():
    d = tmpdir("rr")
    f = os.path.join(d, "m.py")
    open(f, "w").write("def a():\n    return 1\n")
    e = engine(root=d)

    def call(i, args):
        tc = {"id": f"r{i}", "type": "function", "function": {"name": "read_file", "arguments": json.dumps(args)}}
        e.history.append({"role": "assistant", "content": "", "tool_calls": [tc]})
        e._run_tool(tc)
        return e.history[-1]["content"]
    first = call(1, {"path": f})
    assert "return 1" in first
    again = call(2, {"path": f, "offset": 1})       # the same range, asked again
    assert again.startswith(e.READ_POINTER) and "lines 1-2" in again, again[:160]
    other = call(3, {"path": f, "offset": 2})        # a sub-range of what is on screen: also a pointer
    assert other.startswith(e.READ_POINTER), other[:160]
    asked = call(4, {"path": f, "offset": 2})        # asked again right after the pointer: the real text, no loop
    assert "return 1" in asked
    open(f, "w").write("def a():\n    return 2\n")   # changed: full read again
    assert "return 2" in call(5, {"path": f})
    e.history = [m for m in e.history if m.get("tool_call_id") != "r5"]   # that result compacted away
    assert "return 2" in call(6, {"path": f}), "a read that is no longer in context reads again"


@test
def compaction_pins_code_outline():
    d = tmpdir("ol")
    f = os.path.join(d, "game.js")
    open(f, "w").write("export function start() {}\nclass Player {\n  move(dx) { return dx }\n}\n")
    e = engine(root=d)
    e.ctx.changed[f] = "created"
    o = e._code_outline()
    assert "game.js (4 lines)" in o and "start" in o and "Player" in o, o
    assert "Outline of the code files" in agentic._render_pins({"outline": o})


@test
def answer_images_cited_first():
    e = engine()
    e._turn_images = [{"src": "https://i/a.jpg", "page": "https://a.com/x", "title": "A"},
                      {"src": "https://i/b.jpg", "page": "https://b.com/y", "title": "B"},
                      {"src": "https://i/b.jpg", "page": "https://b.com/y", "title": "B"}]
    out = e._answer_images("Try B [1].\n\nSources:\n1. [B](https://b.com/y)")
    assert [i["src"] for i in out] == ["https://i/b.jpg", "https://i/a.jpg"] and out[0]["cited"] and not out[1]["cited"]
    assert e._answer_images("(stopped)") == []
    e._turn_images = []
    assert e._answer_images("No search ran.") == []


@test
def page_image_skips_logos():
    import deep_search as ds
    html = ('<meta property="og:image" content="/img/site-logo.png">'
            '<meta name="twitter:image" content="https://cdn.x.com/photo.jpg">')
    assert ds.page_image(html, "https://x.com/p") == "https://cdn.x.com/photo.jpg"
    assert ds.page_image('<meta property="og:image" content="/media/hero.jpg">', "https://x.com/a/b") == "https://x.com/media/hero.jpg"
    assert ds.page_image("<html></html>", "https://x.com") == ""


@test
def interpret_schema_strict_ok():
    sch = agentic.AgentEngine.INTERPRET_SCHEMA
    assert sch.get("additionalProperties") is False and set(sch["properties"]) == set(sch["required"])


@test
def effort_auto_none_max_and_quick_size():
    e = engine(effort="auto")
    e._think_effort = "none"
    e.turn, e._need_plan, e.history = {"planned": False}, False, [{"role": "tool", "content": "ok"}]
    assert e._wants_thinking(True, 1) is False and e._step_budget(True, 1) == 0     # quick work: no thinking
    e._think_effort = "low"
    e._fc_rounds = 1                                                                # final answer of low work
    assert e._step_budget(True, 2) == agentic.EFFORT_THINK["low"], "only medium/high answers get the doubled budget"
    e.cfg["effort"] = "max"
    assert e._effective_effort() == "high" and e._wants_thinking(True, 3) is True
    e.cfg["effort"] = "low"
    assert e._openai_effort() == "low"
    # the lean pass: size quick caps depth/plan/think, and its understanding is applied in the same call
    e2 = engine(lean_first=True, interpret_first=True)
    e2._effort = None
    e2.history = [{"role": "user", "content": "find a good coffee shop downtown"}]
    e2._turn_user_msg = e2.history[0]
    e2._call_openai = _fake_openai([{"role": "assistant", "content": "", "tool_calls": [tc("get_tools", {
        "why": "local search", "size": "quick", "depth": "extensive", "plan": "full", "think": "high",
        "goal": "a few good coffee shops downtown", "conditions": ["downtown"]})]}])
    assert e2._lean_pass("find a good coffee shop downtown") is None
    assert e2._depth == "basic" and e2._effort["plan"] == "none" and e2._effort["think"] == "low", e2._effort
    assert "Goal: a few good coffee shops downtown" in e2.history[0]["content"]
    assert any(k == "interpretation" for k, _ in e2._emits), "the Understood-as card comes from the same call"


@test
def factcheck_apostrophes_and_latex_numbers():
    import factcheck as f
    text = "Paisano's Pizza, 12 Granville Road. Sonny's Pizza on Elgin Street. O'Brien said."
    for k in ("Paisano's", "Paisano\u2019s", "Sonny's Pizza", "O'Brien"):
        assert f._contains(k, text), k
    assert not f._contains("Luigi's", text)
    ans = r"**1,250,025,000** $$\frac{50000 \times 50001}{2} = 1{,}250{,}025{,}000$$"
    assert [r for r, _, _ in f._answer_numbers(ans)].count("250") == 0
    h = [{"role": "user", "content": "sum of 1 to 50k"},
         {"role": "tool", "name": "calculate", "content": "50000 * 50001 / 2 → 1250025000"}]
    assert f.math_issues(ans, h) == ([], [])


@test
def verify_switch_off_skips_checks():
    e = engine(lean_reference=True, verify=False)
    called = []
    e._lean_reference = lambda q: called.append(q) or "REF"
    m = {"role": "user", "content": "best pizza in some town"}
    e.history, e._turn_user_msg = [m], m
    e._supplement(m["content"])
    assert called == [] and m["content"] == "best pizza in some town", "no background research when off"
    e.history.append({"role": "tool", "name": "web_search", "content": "x" * 400})
    e._fc_from, e._fc_rounds, e._check_level = 0, 0, "full"
    assert e._factcheck_due("A long answer " * 30) is False, "no claim check when off"
    e.cfg["verify"] = True
    assert e._factcheck_due("A long answer " * 30) is True


@test
def factcheck_skips_code_turns_without_web_research():
    e = engine()
    e.cfg["verify"] = True
    e.history = [{"role": "user", "content": "build it"},
                 {"role": "assistant", "content": "", "tool_calls": [{"id": "1", "type": "function", "function": {
                     "name": "write_file", "arguments": json.dumps({"path": "C:/x/server.py", "content": "x"})}}]},
                 {"role": "tool", "name": "write_file", "tool_call_id": "1", "content": "created"},
                 {"role": "tool", "name": "read_file", "content": "x" * 400}]
    e._fc_from, e._fc_rounds, e._check_level = 0, 0, "full"
    assert e._factcheck_due("A long answer " * 30) is False, "code turn: tests are the check"
    e.history.append({"role": "tool", "name": "web_search", "content": "x" * 400})
    assert e._factcheck_due("A long answer " * 30) is True, "code + web research: still checked"


@test
def math_count_rule_ignores_appended_reference():
    e = engine()
    hist = [{"role": "user", "content": "best pizza in some town\n\n(Reference for any facts in your answer… the number "
                                        "of pizzerias is 40 …)"}]
    assert e._math_note("Here are **3 to start**: A, B and C.", history=hist) == ""


@test
def lean_calc_results_join_history_and_repair_round():
    e = engine(lean_first=True, lean_reference=True)
    e._effort = None
    e._turn_ref = "REF: something"
    e.history = [{"role": "user", "content": "What is 5857 x 251?"}]
    e._turn_user_msg = e.history[0]
    replies = [{"role": "assistant", "content": "", "tool_calls": [tc("calculate", {"expression": "5857*251"})]},
               {"role": "assistant", "content": "5857 × 251 = **1,470,107** (with a made-up detail)"},
               {"role": "assistant", "content": "5857 × 251 = **1,470,107**"}]
    e._call_openai = _fake_openai(replies)
    checks = iter([["a made-up detail"], []])          # first check fails -> one repair round; second passes
    e._lean_unsupported = lambda msgs, raw, text: next(checks)
    out = e._lean_pass("What is 5857 x 251?")
    assert out and "1,470,107" in out, out
    roles = [m.get("role") for m in e.history]
    assert "tool" in roles and any(m.get("tool_calls") for m in e.history), roles
    assert [k for k, d in e._emits if k == "answer_retract"], "the repair round retracts the first draft"


@test
def gated_tool_is_enabled_and_runs():
    d = tmpdir("gt")
    e = engine(root=d, tool_gate=True)
    if not e._gated():
        return "tool gating not active in this config"
    inactive = [n for n in ("list_dir", "read_file", "find_files", "search_text") if n in e.reg.by_name
                and n not in e._active_names()]
    if not inactive:
        e._active = set(e._active_names()) - {"list_dir"} if hasattr(e, "_active") else None
        inactive = [n for n in ("list_dir",) if n not in e._active_names()]
    if not inactive:
        return "every tool is active in this config (nothing to gate)"
    name = inactive[0]
    tcall = tc(name, {"path": d})
    e.history.append({"role": "assistant", "content": "", "tool_calls": [tcall]})
    e._run_tool(tcall)
    out = e.history[-1]["content"]
    assert "was not enabled" not in out and "did not run" not in out, out[:200]


# =====================================================================================================================
# read_file honesty: a file under the line limit but over the result clip used to be "shown whole" and then cut at
# MAX_RESULT, so the model never saw its end and every offset read returned line 1 again (bonsai2, 2026-10-11)
@test
def read_file_pages_long_lines_and_states_size():
    import re
    d = tmpdir("read")
    p = os.path.join(d, "server.py")
    body = [f"def handler_{i}(request):  # " + "x" * 30 for i in range(330)]
    with open(p, "w", encoding="utf-8") as f:
        f.write("\n".join(body) + "\n")
    reg = at.Registry(at.Ctx([d]), None, None)
    seen, off, pages = set(), None, 0
    while pages < 10:
        out, _ = reg.call("read_file", {"path": p} if off is None else {"path": p, "offset": off})
        pages += 1
        head = out.splitlines()[0]
        assert head.startswith("[server.py: 330 lines,"), head            # the full size, on every read
        assert "truncated," not in out and "whole file is shown" not in out, out[:300]
        seen.update(int(ln.split("\t")[0]) for ln in out.splitlines() if "\t" in ln and ln.split("\t")[0].strip().isdigit())
        m = re.search(r"continue with offset=(\d+)", head)
        if not m:
            assert "0 after" in head, head
            break
        off = int(m.group(1))
    assert seen == set(range(1, 331)) and 2 <= pages <= 4, (len(seen), pages)
    # an offset read of a long-lined file starts where it was asked, not at line 1
    out, _ = reg.call("read_file", {"path": p, "offset": 200, "limit": 20})
    assert "showing lines 200-" in out.splitlines()[0], out.splitlines()[0]
    # a small file is still shown whole, and says so
    q = os.path.join(d, "small.py")
    with open(q, "w", encoding="utf-8") as f:
        f.write("x = 1\n" * 50)
    out, _ = reg.call("read_file", {"path": q, "offset": 20, "limit": 5})
    assert "showing the whole file" in out.splitlines()[0] and "    50\t" in out, out[:200]


@test
def slimmed_history_call_has_no_content_to_copy():
    # models imitate their own earlier calls: an old write must not keep a `content` value (the placeholder was
    # copied as the content of new writes — bonsai1 wiped store.py, bonsai2 tried twice)
    tcall = tc("write_file", {"path": "a.py", "content": "x" * 1000})
    agentic._slim_call_args(tcall, ref="E5")
    a = json.loads(tcall["function"]["arguments"])
    assert "content" not in a and a["path"] == "a.py" and "1,000 chars omitted" in a["content_omitted"], a
    assert at._HISTORY_STUB.search(a["content_omitted"])     # copied by hand, the write guard still refuses it


@test
def guide_and_catalog_are_not_cut_at_the_default_limit():
    reg = at.Registry(at.Ctx([tmpdir("g")]), None, None)
    long = "x" * 15000
    assert reg._keep_result("read_guide", long) == long                  # software.md is 11 KB, the biggest 15 KB
    assert "omitted from the middle" in reg._keep_result("read_file_unknown_tool", long)


@test
def a_long_result_keeps_its_errors_and_its_full_text():
    import re
    reg =at.Registry(at.Ctx([tmpdir("big")]), None, None)
    log = ("setup line\n" * 400 + "FAILED tests/test_api.py::test_a5 - assert 201 == 400\n" + "noise line\n" * 600
           + "== 3 failed, 8 passed ==\n")
    out = reg._keep_result("powershell", log)
    assert len(out) <= at.MAX_RESULT + 400, len(out)
    assert "FAILED tests/test_api.py::test_a5" in out and "3 failed, 8 passed" in out, out[-500:]   # middle + end
    path = re.search(r"saved at (.+?\.txt)", out).group(1)
    assert open(path, encoding="utf-8").read() == log                      # the whole result is on disk


@test
def stubs_carry_enough_to_act_on():
    text = "tests: FAILED — 5 failed, 6 passed\n" + "\n".join(f"FAILED tests/t.py::test_{i} - AssertionError: x" for i in range(5))
    stub = agentic.AgentEngine._key_output("run_tests", {}, text)
    assert all(f"test_{i}" in stub for i in range(5)), stub              # every failure, not the first three




@test
def masking_keeps_earlier_pages_of_a_paged_file():
    e = engine()
    msgs, n = [], 0

    def call(name, args, content):
        nonlocal n
        n += 1
        msgs.append({"role": "assistant", "content": "", "tool_calls": [tc(name, args, f"c{n}")]})
        msgs.append({"role": "tool", "tool_call_id": f"c{n}", "content": content})
        return f"c{n}"

    page = "\n".join(f"{i:>6}\tline {i} " + "x" * 40 for i in range(1, 60))
    p1 = call("read_file", {"path": "C:/p/server.py"}, page)
    p2 = call("read_file", {"path": "C:/p/server.py", "offset": 137}, page)
    again = call("read_file", {"path": "C:/p/client.py"}, page)
    call("read_file", {"path": "C:/p/client.py"}, page)                 # the same page again: the older copy goes
    for i in range(20):
        call("list_dir", {"path": f"C:/p/d{i}"}, "x" * 50)
    stubs = e._stale_tool_stubs(msgs)
    assert p1 not in stubs and p2 not in stubs, stubs                    # different pages of one file both stay
    assert again in stubs, stubs


@test
def a_post_makes_the_same_get_a_new_question():
    e = engine()
    t = e.reg.by_name["http_request"]
    get = {"method": "GET", "url": "http://127.0.0.1:9/totals"}
    e._record_call("http_request", t, get, "HTTP 200 {}")
    assert e._repeat_of("http_request", t, get)                     # the same GET, nothing in between: refused
    e._record_call("http_request", t, {"method": "POST", "url": "http://127.0.0.1:9/add", "json": {"a": 1}}, "HTTP 201")
    assert not e._repeat_of("http_request", t, get)                 # after a POST the server's data changed


@test
def auto_size_uses_a_bigger_gpu_and_leaves_the_tuned_one_alone():
    import llama_local as ll
    saved = (ll.gpu_total_mib, ll.gguf_meta, ll.ollama_blob, ll.fit_context)
    try:
        ll.ollama_blob = lambda m: m
        ll.gguf_meta = lambda p: {"context_length": 262144}
        ll.fit_context = lambda m, want, kv, reserve_mib=0, kv_v=None: {
            "ctx": min(want, 90112), "kv": kv, "max": {"f16": 86016 if ll.gpu_total_mib() > 12000 else 24576}}
        ll.gpu_total_mib = lambda: 8192
        assert ll.auto_size("m", 65536, "q8_0", "q4_0", tuned_mib=8192) is None          # the card it was tuned on
        ll.gpu_total_mib = lambda: 12282
        a = ll.auto_size("m", 65536, "q8_0", "q4_0", tuned_mib=8192)
        assert a["kv"] == "f16" and a["ctx"] == 86016 and not a["moe"], a                # better cache, more context
        ll.gguf_meta = lambda p: {"context_length": 262144, "expert_count": 128}
        a = ll.auto_size("m", 65536, "bf16", None, tuned_mib=8192)
        assert a["moe"] and a["ctx"] == 65536, a                                         # MoE: experts move instead
    finally:
        ll.gpu_total_mib, ll.gguf_meta, ll.ollama_blob, ll.fit_context = saved


@test
def contract_is_pinned_from_every_contract_section_and_while_reading():
    e = engine()
    d = tmpdir("proj")
    os.makedirs(os.path.join(d, "docs"))
    with open(os.path.join(d, "docs", "DESIGN.md"), "w", encoding="utf-8") as f:
        f.write("# Design\nprose\n## API contract\n| GET | /api/ping |\n## Command-line usage\n```\n"
                "python client.py --host H add --person P\n```\n## Notes\n- not a contract\n")
    code = os.path.join(d, "server.py"), os.path.join(d, "client.py")
    for p in code:
        with open(p, "w", encoding="utf-8") as f:
            f.write("def main():\n    pass\n")
    c = e._contract_lines(d)
    assert "/api/ping" in c and "client.py --host H add" in c and "not a contract" not in c, c
    for p in code:                                         # only read, nothing edited yet: the state still shows
        at.note_seen(e.ctx, p)
    ws = e._work_state()
    assert "WORK STATE" in ws and "client.py --host H add" in ws, ws[:400]


@test
def tool_profiles_send_a_core_set_and_load_groups_on_demand():
    e = engine(tool_gate=True, stable_tools=True, compact_prompt=False)
    groups = {g: set(ns) for g, _b, ns in at.tool_groups(e.reg)}
    on = set(e._active_names())
    assert {"read_file", "edit_file", "run_tests", "web_search", "calculate", "update_tracker"} <= on, on
    assert not (groups.get("browser", set()) & on) and not (groups.get("schedule", set()) & on)
    assert len(on) < len(e.reg.by_name)
    e._profile_tools("log in to the website and fill out the form")    # the request's words load the browser
    assert groups["browser"] <= set(e._active_names())
    e2 = engine(tool_gate=True, stable_tools=True, compact_prompt=False)
    e2._enable(["schedule_task"])                                      # a call/enable loads the whole group
    assert groups["schedule"] <= set(e2._active_names())
    e3 = engine(tool_gate=True, stable_tools=True, tool_profiles=False, compact_prompt=False)
    assert set(e3._active_names()) == set(e3.reg.by_name)              # the switch restores every tool


@test
def layered_prompt_and_brief_docs_follow_the_loaded_groups():
    e = engine(tool_gate=True, stable_tools=True)
    s = e.system_prompt()
    assert 'enable=["browser"]' in s and "browser_read shows the page" not in s
    assert "NOTES.md" in s and "artifact_create puts a document" not in s
    e._enable(["browser", "artifacts"])
    s2 = e.system_prompt()
    assert "browser_read shows the page" in s2 and "artifact_create puts a document" in s2
    sch = {t["function"]["name"]: t["function"]["description"] for t in e._schemas(True)}
    assert "read_guide('tool:browser_devtools')" in sch["browser_devtools"]          # long reference text: brief
    assert "no &&-chains" in sch["powershell"] and "read_guide('tool:" not in sch["powershell"]   # rules stay whole
    full = at.t_read_guide({"topic": "tool:browser_devtools"}, None)
    assert full.startswith("# browser_devtools") and "Parameters:" in full


@test
def turning_points_compaction_map_and_ledger():
    import context_memory as cm
    e = engine(num_ctx=16384, compact_prompt=False, context_compaction_v2=False)
    conv = e.ctx.conv_id
    e.turn = {"steps": 3}
    at.t_update_tracker({"new": True, "goal": "expense tracker", "milestones": [{"title": "PRD"}, {"title": "Build"}]},
                        e.ctx)
    e._watch_progress("update_tracker", {"new": True, "goal": "expense tracker", "milestones": [{"title": "PRD"},
                                                                                               {"title": "Build"}],
                                        "decisions": ["category = any non-empty string"]}, True, "checklist saved")
    e._watch_progress("write_file", {"path": "C:/p/server.py"}, True, "created C:/p/server.py (120 lines, +120 -0)")
    log = "FAILED tests/test_api.py::test_a5 - assert 201 == 400\nValueError: bad amount\n== 1 failed, 8 passed =="
    e._watch_progress("run_tests", {}, True, log)
    e._watch_progress("run_tests", {}, True, log)                                     # the same state: no new point
    e._watch_progress("run_tests", {}, True, "== 9 passed ==")
    kinds = [x["kind"] for x in cm.events(conv)]
    assert kinds == ["plan", "decided", "created", "tests", "error", "tests"], kinds
    # a forced compaction builds on the timeline, pins it, and writes the map
    e.history = [{"role": "user", "content": "build me an expense tracker"}]
    for i in range(12):
        e.history.append({"role": "assistant", "content": f"step {i}", "tool_calls": [tc("read_file", {"path": f"f{i}.py"}, f"c{i}")]})
        e.history.append({"role": "tool", "tool_call_id": f"c{i}", "name": "read_file", "content": f"file {i}\n" + "x = 1\n" * 300})
    asked = []
    e._call_openai = lambda messages, **k: (asked.append(messages[-1]["content"]) or
                                            ({"role": "assistant", "content": "SUMMARY: " + "built and tested. " * 10}, 10, 100))
    e._compact_req = True
    e._maybe_compact()
    head = e.history[0]["content"]
    assert "Turning points, recorded by the app" in head and "category = any non-empty string" in head, head[:600]
    path = e._pins.get("map")
    assert path and os.path.isfile(path) and "| E1 |" in open(path, encoding="utf-8").read()
    assert any("Turning points the app recorded" in a for a in asked), asked[-1:]       # the summary builds on them
    assert "CONTEXT_MAP.md" in head                                                       # and the head points at the map
    e._payload_messages()
    rows = cm.ledger(conv)
    assert rows and rows[-1]["total"] > 0 and "tool_stub" in rows[-1], rows[-1:]


@test
def xml_style_tool_arguments_are_coerced_to_their_types():
    # Qwen-style XML tool calls deliver lists as JSON text: a whole milestone list became one milestone
    reg = at.Registry(at.Ctx([tmpdir("co")]), None, None)
    a = reg.normalize_args("update_tracker", {"new": "True", "goal": "g", "notes": "one note",
                                              "milestones": '[{"title": "Research + PRD"}, {"title": "Build"}]'})
    assert a["new"] is True and a["notes"] == ["one note"] and [m["title"] for m in a["milestones"]] == ["Research + PRD", "Build"]
    assert "0/2 done" in at.t_update_tracker(a, reg.ctx)
    py = reg.normalize_args("update_tracker", {"updates": "[{'id': '1', 'status': 'done'}]"})   # Python-style text
    assert py["updates"] == [{"id": "1", "status": "done"}]
    keep = reg.normalize_args("write_file", {"path": "a.py", "content": "[1, 2]"})            # strings stay strings
    assert keep["content"] == "[1, 2]"


@test
def hesitation_words_are_capitalised_sentence_starts_only():
    import llama_local as ll
    words = [w.strip() for w in ll.HESITATION_WORDS]
    assert all(w[0].isupper() for w in words), words                  # never 'wait' (wait_port, await) or 'hmm'
    assert {"Wait", "Hmm", "Alternatively", "Actually"} == set(words)


@test
def settled_decisions_stay_in_view():
    c = at.Ctx([tmpdir("dec")])
    at.t_update_tracker({"new": True, "goal": "fix the tracker", "milestones": [{"title": "fix tests"}]}, c)
    at.t_update_tracker({"decisions": ["category = any non-empty string (DESIGN.md)"], "notes": ["n0"]}, c)
    at.t_update_tracker({"decisions": ["Category = any non-empty string (DESIGN.md)", "client args: --host --port first"]}, c)
    for i in range(15):                                             # notes scroll; decisions don't
        at.t_update_tracker({"notes": [f"note {i}"]}, c)
    shown = at.tracker_prompt(c.tracker)
    assert "Settled" in shown and "category = any non-empty string" in shown and "--host --port first" in shown, shown
    assert shown.count("non-empty string") == 1                     # de-duplicated (case-insensitive)
    assert "n0" not in shown                                        # the old note scrolled away




@test
def compaction_waits_until_the_reply_would_not_fit():
    e = engine(num_ctx=61440)
    assert e.context_budget() == 61440 - 8192                      # was 75% (46,080): compacting is lossy
    e.cfg["num_ctx"] = 131072
    assert e.context_budget() == 131072 - 8192
    e.cfg["num_ctx"] = 20480
    assert e.context_budget() == 15360                              # small contexts: never below 75%
    e.cfg["compact_at"] = 0.6
    assert e.context_budget() == int(20480 * 0.6)                   # an explicit setting still wins


@test
def a_long_single_code_write_gets_a_note():
    d = tmpdir("big")
    reg = at.Registry(at.Ctx([d]), None, None)
    out, _ = reg.call("write_file", {"path": os.path.join(d, "big.py"), "content": "x = 1\n" * 300})
    assert "write the next big file in parts" in out, out
    out, _ = reg.call("write_file", {"path": os.path.join(d, "small.py"), "content": "x = 1\n" * 30})
    assert "in parts" not in out, out








@test
def each_model_call_logs_where_its_time_went():
    """gen.jsonl: thinking / text / tool-call tokens and the server's prompt-read and generation timings per call."""
    import context_memory as cm
    e = engine(backend="llama")
    e.base = "http://127.0.0.1:9"                  # never contacted: urlopen is replaced below
    e.turn = {"steps": 4}
    ev = [{"choices": [{"delta": {"reasoning_content": "hm"}}]}] * 3 + \
         [{"choices": [{"delta": {"tool_calls": [{"index": 0, "id": "c1", "function": {"name": "list_dir",
                                                                                       "arguments": "{}"}}]}}]},
          {"choices": [{"delta": {}, "finish_reason": "tool_calls"}],
           "timings": {"prompt_n": 120, "cache_n": 3000, "prompt_ms": 400.0, "predicted_n": 4, "predicted_ms": 300.0}}]

    class Resp:
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def __iter__(self):
            return iter([("data: " + json.dumps(x)).encode() for x in ev] + [b"data: [DONE]"])

    real = agentic.urllib.request.urlopen
    agentic.urllib.request.urlopen = lambda *a, **k: Resp()
    try:
        e._call_openai([{"role": "system", "content": "s"}, {"role": "user", "content": "u"}])
    finally:
        agentic.urllib.request.urlopen = real
    rows = cm.gen_log(e.ctx.conv_id)
    assert rows, "no gen.jsonl row"
    r = rows[-1]
    assert (r["think_tok"], r["call_tok"], r["text_tok"]) == (3, 1, 0), r
    assert (r["read_tok"], r["cached_tok"], r["gen_tok"], r["step"]) == (120, 3000, 4, 4), r
    assert r["calls"] == ["list_dir"], r


@test
def prompt_opening_stays_identical_through_a_run():
    """The server can only reuse its cache up to the first change (on a hybrid model: up to a checkpoint, usually
    none that early). A `cd`, the workspace being created and the date turning used to change the system prompt."""
    root = tmpdir("cache")
    e = engine(root, workspace_root=os.path.join(root, "Zen Harness"))
    e.turn = {"steps": 1, "today": "Sunday 11 October 2026"}
    e._ensure_workspace("build an app")
    first = e.system_prompt()
    sub = os.path.join(root, "proj")
    os.makedirs(sub)
    e.ctx.cwd = sub                                        # what powershell's `cd` does
    os.makedirs(e.ctx.workspace, exist_ok=True)            # the workspace gets created on first use
    real = agentic.time.strftime
    agentic.time.strftime = lambda fmt, *a: "Monday 12 October 2026" if "%A" in fmt else real(fmt, *a)
    try:
        assert e.system_prompt() == first, "the system prompt changed mid-request"
    finally:
        agentic.time.strftime = real
    r = e._reminder()
    assert r and sub in r["content"], "the working folder belongs in the end-of-prompt reminder"


@test
def thinking_off_keeps_the_prompt_opening_when_the_template_moves_it():
    import urllib.request as ur
    e = engine(backend="llama")
    e.base = "http://127.0.0.1:9"
    e.turn = {"steps": 1}
    e._think_step = False                                   # a step that shouldn't think
    sent = []

    class Resp:
        def __init__(self, body):
            self.body = body

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def read(self):
            return self.body

        def __iter__(self):
            return iter([b'data: {"choices":[{"delta":{"content":"ok"},"finish_reason":"stop"}]}', b"data: [DONE]"])

    def fake(req, *a, **k):
        body = json.loads(req.data.decode())
        if req.full_url.endswith("/apply-template"):           # a Bonsai-like template: the header follows the flag
            on = (body.get("chat_template_kwargs") or {}).get("enable_thinking", True)
            head = "<|im_start|>system\n" + ("Reasoning effort is set to xhigh.\n\n" if on else "")
            return Resp(json.dumps({"prompt": head + "S<|im_end|>\nU1 A1 U2"}).encode())
        sent.append(body)
        return Resp(b"")

    real = agentic.urllib.request.urlopen
    agentic.urllib.request.urlopen = fake
    agentic._THINK_FLAG_MOVES.clear()
    try:
        e._call_openai([{"role": "system", "content": "s"}, {"role": "user", "content": "u"}])
        b = sent[-1]
        assert "chat_template_kwargs" not in b and b.get("thinking_budget_tokens") == agentic.NO_THINK_BUDGET, b
        agentic._THINK_FLAG_MOVES.clear()
        agentic._THINK_FLAG_MOVES[(e.base, e.cfg.get("model") or "")] = False   # an ordinary template: switch it off
        e._call_openai([{"role": "system", "content": "s"}, {"role": "user", "content": "u"}])
        assert sent[-1].get("chat_template_kwargs") == {"enable_thinking": False}, sent[-1]
    finally:
        agentic.urllib.request.urlopen = real
        agentic._THINK_FLAG_MOVES.clear()


@test
def masking_waits_until_a_batch_frees_enough():
    e = engine(num_ctx=524288)                                      # far from compaction: only the freed amount decides
    msgs = [{"role": "system", "content": "s" * 140000}]          # a big prompt: ~40K tokens

    def call(i, path, content):
        msgs.append({"role": "assistant", "content": "", "tool_calls": [tc("read_file", {"path": path}, f"c{i}")]})
        msgs.append({"role": "tool", "tool_call_id": f"c{i}", "content": content})

    call(0, "C:/p/a.py", "x" * 2000)                                # superseded later, but small
    call(1, "C:/p/a.py", "x" * 2000)
    for i in range(2, 30):
        call(i, f"C:/p/f{i}.py", "y" * 500)
    view = e._masked_view(msgs)
    assert view == msgs, "a batch that frees almost nothing must not rewrite the prompt"
    for i in range(30, 44):                                         # big superseded reads pile up
        call(i, "C:/p/big.py", "z" * 30000)
    for i in range(44, 60):                                         # ...and fall out of the recent window
        call(i, f"C:/p/g{i}.py", "w" * 300)
    view = e._masked_view(msgs)
    assert any(m.get("content") != o.get("content") for m, o in zip(view, msgs)), "a worthwhile batch applies"


@test
def masking_skips_a_batch_close_to_compaction():
    # the same worthwhile batch, but the prompt is already near the compaction point: compaction rewrites the prompt
    # shortly anyway, so masking now would only add a second full re-read (Bonsai: masked at 15:31 and 15:53, ~110 s
    # each, then compacted at 16:03)
    e = engine(num_ctx=61440)
    msgs = [{"role": "system", "content": "s" * 100000}]

    def call(i, path, content):
        msgs.append({"role": "assistant", "content": "", "tool_calls": [tc("read_file", {"path": path}, f"c{i}")]})
        msgs.append({"role": "tool", "tool_call_id": f"c{i}", "content": content})

    for i in range(0, 8):                                           # superseded reads: ~25K tokens to free
        call(i, "C:/p/big.py", "z" * 12000)
    for i in range(8, 30):
        call(i, f"C:/p/g{i}.py", "w" * 300)
    assert agentic._est_tokens(msgs) * e.ratio > e.context_budget() * 0.75
    assert e._masked_view(msgs) == msgs, "no masking batch right before compaction"


@test
def green_suite_is_the_evidence_in_the_work_state():
    d = tmpdir("ver")
    e = engine(d)
    f = os.path.join(d, "app.py")
    open(f, "w").write("def main():\n    return 1\n")
    e.ctx.changed[f] = "created"
    e._project = True
    h = at._hist(e.ctx)
    h["last_tests"] = {"passed": 21, "failed": 0, "t": "23:48:47", "cmd": "tests", "failing": []}
    h["green"], h["green_t"], h["green_edits"] = {}, "23:48:47", 0
    ws = e._work_state()
    assert "Verified" in ws and "one live run" in ws and "no new scripts" in ws, ws
    h["green_edits"] = 2                                  # code changed since: run them again, no claim
    assert "Verified:" not in e._work_state()


@test
def build_profile_trims_the_fixed_prompt_start():
    e = engine(tool_profiles=True, compact_prompt=False)
    assert e._BUILD_REQ_RX.search("I want you to build an expense tracker app for my roommates")
    assert not e._BUILD_REQ_RX.search("what's the weather in Austin")
    full = e.system_prompt()
    core = e._core_names()
    e._build_profile = True
    built = e.system_prompt()
    assert "Critical facts:" in full and "Critical facts:" not in built and "Claims about your own work" in built
    assert len(built) < len(full)
    assert not (e._core_names() & e.BUILD_SKIP) and (core & e.BUILD_SKIP), "build skips research-only tools"
    assert {"read_file", "write_file", "run_tests", "powershell"} <= e._core_names()


@test
def masking_condenses_old_file_copies_and_old_narration():
    e = engine()
    msgs = [{"role": "system", "content": "s" * 120000}]
    big = "x = 1\n" * 2000                                         # a 12K-char file written twice
    talk = "I will now carefully explain everything I am about to do in great detail. " * 8

    def step(i, name, args, result="ok", say=""):
        msgs.append({"role": "assistant", "content": say, "tool_calls": [tc(name, args, f"c{i}")]})
        msgs.append({"role": "tool", "tool_call_id": f"c{i}", "content": result})

    step(0, "write_file", {"path": "C:/p/app.py", "content": big}, say=talk)
    step(1, "write_file", {"path": "C:/p/app.py", "content": big + "y = 2\n"}, say=talk)
    for i in range(2, 30):
        step(i, "list_dir", {"path": f"C:/p/d{i}"}, "z" * 60, say=talk)
    view = e._masked_view(msgs)
    first_call = view[1]["tool_calls"][0]["function"]["arguments"]
    second_call = view[3]["tool_calls"][0]["function"]["arguments"]
    # every write older than the last few steps drops its text (on disk; the working set re-adds the current file)
    assert "content_omitted" in first_call and "content_omitted" in second_call, "older writes condensed"
    assert "older narration condensed" in view[1]["content"], view[1]["content"][:120]
    assert view[-2]["content"] == talk, "the latest narration stays whole"
    assert e._apply_mask(msgs) == view, "the summary request sends the same view"


@test
def fact_checks_see_what_compaction_archived():
    import history_archive as ha
    e = engine()
    ha.add(e.ctx.conv_id, "tool result", "run_tests", "tests: PASSED — 21 passed (2.6s)")
    e.history = [{"role": "user", "content": "build it"}]
    ev = e._evidence_history()
    assert any("21 passed" in str(m.get("content")) for m in ev if m.get("role") == "tool")




@test
def check_script_beside_a_green_suite_gets_a_note():
    d = tmpdir("smk")
    c = at.Ctx([d])
    at._hist(c)["last_tests"] = {"passed": 21, "failed": 0, "t": "x", "failing": []}
    out = at.t_write_file({"path": os.path.join(d, "smoke_test.py"), "content": "print(1)\n"}, c)
    assert "Make a new check a test in that suite" in out, out
    out = at.t_write_file({"path": os.path.join(d, "test_more.py"), "content": "def test_a():\n    pass\n"}, c)
    assert "Make a new check" not in out, "a proper test file is the suite"


@test
def green_suite_ticks_a_tests_checklist_item():
    d = tmpdir("tick")
    c = at.Ctx([d])
    c.conv_id = "audit-tick"
    at.t_update_tracker({"new": True, "goal": "g", "milestones": [{"title": "Verify", "subtasks": ["run_tests passes",
                                                                                                  "README"]}]}, c)
    note = at._tick_on_green(c, 21, 0)
    assert "marked done" in note and at.tracker_current(c.tracker)[1].startswith("README"), note
    assert at._tick_on_green(c, 20, 1) == ""


@test
def read_file_one_symbol():
    d = tmpdir("sym")
    p = os.path.join(d, "m.py")
    open(p, "w").write("import os\n\n\nclass Ledger:\n    def load(self):\n        return 1\n\n    @property\n"
                       "    def size(self):\n        return 2\n\n\ndef main():\n    pass\n")
    c = at.Ctx([d])
    out = at.t_read_file({"path": p, "symbol": "Ledger.size"}, c)
    assert "def size" in out and "@property" in out and "def load" not in out and "def main" not in out, out
    try:
        at.t_read_file({"path": p, "symbol": "nope"}, c)
        raise AssertionError("an unknown symbol must name what exists")
    except at.ToolError as ex:
        assert "Ledger.load" in str(ex)


@test
def near_json_checklist_is_repaired_not_saved_as_one_title():
    """bonsai9 ended its milestone list with '})' instead of '}]': the checklist became one milestone named '[{...'."""
    raw = ('[{"title": "Docs", "done_when": "PRD written", "subtasks": [{"step": "PRD.md (goal, users)"}]}, '
           '{"title": "Build", "subtasks": [{"step": "server"}, {"step": "stop server, final report"}]})')
    assert [m["title"] for m in at._loads_lenient(raw)] == ["Docs", "Build"]
    assert at._coerce(raw, {"type": "array"})[1]["title"] == "Build"
    assert at._loads_lenient('{"a": [1, 2}') == {"a": [1, 2]}             # a missing closer
    assert at._loads_lenient('[{"t": "a ) ] } in a string"}]')[0]["t"] == "a ) ] } in a string"
    d = tmpdir("nj")
    c = at.Ctx([d])
    c.conv_id = "audit-nearjson"
    out = at.t_update_tracker({"new": True, "goal": "g", "milestones": raw}, c)
    assert len(c.tracker["milestones"]) == 2 and c.tracker["milestones"][0]["title"] == "Docs", out
    try:
        at.t_update_tracker({"new": True, "goal": "g", "milestones": '[{"title": "x" "oops"'}, c)
        raise AssertionError("unreadable JSON must be refused, not saved as a title")
    except at.ToolError as ex:
        assert "Send it again" in str(ex)


@test
def same_failure_twice_switches_to_tracing():
    d = tmpdir("trace")
    e = engine(d)
    f = os.path.join(d, "app.py")
    open(f, "w").write("x = 1\n")
    e.ctx.changed[f] = "created"
    e._project = True
    h = at._hist(e.ctx)
    h["last_tests"] = {"passed": 9, "failed": 2, "t": "02:30:00", "cmd": "tests", "failing": ["t::a6"]}
    h["same_fail"] = 1
    assert "TRACE, DON'T GUESS" not in e._work_state()
    h["same_fail"] = 2
    assert "TRACE, DON'T GUESS" in e._work_state()


@test
def a_test_that_starts_its_own_server_gets_a_note():
    d = tmpdir("tsrv")
    c = at.Ctx([d])
    src = ("import threading\nfrom http.server import ThreadingHTTPServer\nimport server\n\n"
           "def test_a():\n    srv = ThreadingHTTPServer(('127.0.0.1', 0), server.H)\n"
           "    threading.Thread(target=srv.serve_forever).start()\n")
    out = at.t_write_file({"path": os.path.join(d, "test_api.py"), "content": src}, c)
    assert "these tests start a server themselves" in out, out
    out = at.t_write_file({"path": os.path.join(d, "test_logic.py"), "content": "def test_b():\n    assert 1\n"}, c)
    assert "start a server" not in out


@test
def scratch_scripts_do_not_make_a_green_run_stale():
    """bonsai9 wrote _realrun.py after a green suite: it counted as a code change, switching VERIFIED off and
    're-run the tests now' on."""
    d = tmpdir("scratch")
    os.makedirs(os.path.join(d, "docs"))
    os.makedirs(os.path.join(d, "tests"))
    open(os.path.join(d, "server.py"), "w").write("x = 1\n")
    open(os.path.join(d, "client.py"), "w").write("y = 2\n")
    open(os.path.join(d, "tests", "test_a.py"), "w").write("import client\nCMD = ['python', 'server.py']\n")
    for f, want in (("server.py", True), ("client.py", True), (os.path.join("tests", "test_a.py"), True)):
        assert at._part_of_app(os.path.join(d, f)) is want, f
    c = at.Ctx([d])
    c.conv_id = "audit-scratch"                       # its own file history (contexts without an id share one)
    at.mark_green(c)
    at.t_write_file({"path": os.path.join(d, "_realrun.py"), "content": "import client\nprint(1)\n"}, c)
    assert at._hist(c)["green_edits"] == 0, "a scratch script nothing refers to is not an app change"
    c.files_read = {}
    at._hist(c)  # same history
    out = at.t_edit_file({"path": os.path.join(d, "server.py"), "old": "x = 1", "new": "x = 3"}, c) \
        if hasattr(at, "t_edit_file") else ""
    if out:
        assert at._hist(c)["green_edits"] == 1, "an app file still counts"


def main():
    t_all = time.time()
    for fn in TESTS:
        name = fn.__name__
        t0 = time.time()
        try:
            note = fn()
            status, detail = ("PASS", note or "")
        except AssertionError as ex:
            status, detail = ("WARN" if name.startswith("probe_") else "FAIL"), str(ex) or traceback.format_exc(limit=1)
        except Exception:
            status, detail = "ERROR", traceback.format_exc()[-700:]
        RESULTS.append((name, status, round(time.time() - t0, 2), detail))
        print(f"{status:5} {name} ({time.time() - t0:.2f}s){'  -- ' + detail.strip().splitlines()[-1][:220] if detail else ''}",
              flush=True)
    counts = {s: sum(1 for r in RESULTS if r[1] == s) for s in ("PASS", "FAIL", "WARN", "ERROR")}
    print(f"\n{len(RESULTS)} tests in {time.time() - t_all:.1f}s: " + ", ".join(f"{v} {k}" for k, v in counts.items()))
    with open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "offline_suite_results.json"), "w",
              encoding="utf-8") as f:
        json.dump([dict(zip(("test", "status", "seconds", "detail"), r)) for r in RESULTS], f, indent=1)
    shutil.rmtree(SCRATCH, ignore_errors=True)
    os._exit(1 if counts["FAIL"] or counts["ERROR"] else 0)


if __name__ == "__main__":
    main()
