"""Create real, professionally styled Office documents for the agent.

  pptx  {"theme": "...", "slides": [{"layout": "title|bullets|cards|two_column|stats|table|chart|timeline|quote|
        section|image|closing", ...}]}  or Markdown (one # heading per slide, - bullets)
  docx  Markdown (# headings, paragraphs, - / 1. lists, | tables |, ```chart blocks, ![](image), > callouts,
        **bold**, *italic*, `code`, ``` code blocks, --- page break)
  xlsx  {"theme": "...", "sheets": [{"name", "title", "rows", "formats", "totals", "charts": [...]}]} or CSV text
  pdf   the same Markdown as docx, laid out as HTML and printed by Edge (doc_pdf.py); Word is the fallback
Themes live in doc_themes.py and are applied the same way to all four formats.
See guides/*.md (read_guide) for the formats the agent is taught.
"""
import csv
import datetime
import io
import json
import os
import re

import doc_themes as dt
from doc_themes import THEMES, mix, paper, theme  # noqa: F401  (re-exported for older callers)


class DocError(Exception):
    pass


# ================================================================== Word
_INLINE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*|_[^_\s][^_]*_|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def _rgb(hexs):
    from docx.shared import RGBColor
    return RGBColor.from_string(hexs)


def _font(target, name=None, size=None, color=None, bold=None, italic=None):
    """Set a run/style font fully: explicit names win over the template's theme-font and theme-colour hooks."""
    from docx.oxml.ns import qn
    from docx.shared import Pt
    f = target.font
    if name:
        f.name = name
        rfonts = target.element.get_or_add_rPr().find(qn("w:rFonts"))
        if rfonts is not None:
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rfonts.attrib.pop(qn(a), None)
            rfonts.set(qn("w:eastAsia"), name)
            rfonts.set(qn("w:cs"), name)
    if size:
        f.size = Pt(size)
    if color:
        f.color.rgb = _rgb(color)
        c = target.element.get_or_add_rPr().find(qn("w:color"))
        if c is not None:
            for a in ("w:themeColor", "w:themeShade", "w:themeTint"):
                c.attrib.pop(qn(a), None)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic


def _runs(par, text, color=None, size=None, link=None, mono="Consolas"):
    for part in _INLINE.split(str(text)):
        if not part:
            continue
        if part[:2] in ("**", "__") and part[-2:] == part[:2] and len(part) > 4:
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part[0] in "*_" and part[-1] == part[0] and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.italic = True
        elif part[0] == "`" and part[-1] == "`":
            r = par.add_run(part[1:-1])
            r.font.name = mono
        elif part.startswith("[") and "](" in part:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            r = par.add_run(m.group(1))
            r.underline = True
            if link and not color:
                r.font.color.rgb = _rgb(link)
        else:
            r = par.add_run(part)
        if color:
            r.font.color.rgb = _rgb(color)
        if size:
            from docx.shared import Pt
            r.font.size = Pt(size)


def _el(tag, **attrs):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


def _shade(cell_or_par, hexs):
    from docx.oxml.ns import qn
    host = cell_or_par._tc.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else cell_or_par._p.get_or_add_pPr()
    for old in host.findall(qn("w:shd")):
        host.remove(old)
    host.append(_el("w:shd", val="clear", color="auto", fill=hexs))


def _par_border(par, **sides):
    """sides: bottom=(size_eighths, colour, space_pt)"""
    from docx.oxml.ns import qn
    ppr = par._p.get_or_add_pPr() if hasattr(par, "_p") else par.element.get_or_add_pPr()
    old = ppr.find(qn("w:pBdr"))
    if old is not None:
        ppr.remove(old)
    bdr = _el("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        if side in sides:
            sz, col, space = sides[side]
            bdr.append(_el("w:" + side, val="single", sz=sz, space=space, color=col))
    ppr.append(bdr)


def _cell_border(cell, **sides):
    from docx.oxml.ns import qn
    tcpr = cell._tc.get_or_add_tcPr()
    old = tcpr.find(qn("w:tcBorders"))
    if old is not None:
        tcpr.remove(old)
    b = _el("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        if side in sides:
            sz, col = sides[side]
            b.append(_el("w:" + side, val="single" if sz else "nil", sz=sz or 0, space=0, color=col or "auto"))
    tcpr.append(b)


def _field(par, code, size=None, color=None):
    run = par.add_run()
    if size or color:
        _font(run, size=size, color=color)
    run._r.append(_el("w:fldChar", fldCharType="begin"))
    instr = _el("w:instrText")
    instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr.text = f" {code} "
    run._r.append(instr)
    run._r.append(_el("w:fldChar", fldCharType="end"))


def _docx_styles(doc, p):
    from docx.shared import Pt
    hfont, bfont = p["fonts"]
    normal = doc.styles["Normal"]
    _font(normal, bfont, 10.5, p["text"])
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2
    spec = {1: (17, p["title"], True), 2: (13, p["title"], True), 3: (11, p["accent"], True), 4: (10.5, p["muted"], True)}
    for lvl, (size, col, bold) in spec.items():
        st = doc.styles[f"Heading {lvl}"]
        _font(st, hfont, size, col, bold=bold, italic=False)
        pf = st.paragraph_format
        pf.space_before = Pt({1: 20, 2: 14, 3: 10, 4: 8}[lvl])
        pf.space_after = Pt(4 if lvl > 1 else 6)
        pf.keep_with_next = True
    _par_border(doc.styles["Heading 1"], bottom=(6, p["line"], 3))
    t = doc.styles["Title"]
    _font(t, hfont, 28, p["title"], bold=True)
    t.paragraph_format.space_after = Pt(2)
    _par_border(t, bottom=(1, "FFFFFF", 0))           # drop the template's blue underline
    st = doc.styles["Subtitle"]
    _font(st, bfont, 13, p["muted"], bold=False, italic=False)
    st.paragraph_format.space_after = Pt(18)
    for name in ("List Bullet", "List Bullet 2", "List Number"):
        doc.styles[name].paragraph_format.space_after = Pt(3)


def _docx_table(doc, rows, align, p):
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt
    ncol = len(rows[0])
    tbl = doc.add_table(rows=len(rows), cols=ncol)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblpr = tbl._tbl.tblPr
    mar = _el("w:tblCellMar")
    for side, w in (("top", 70), ("bottom", 70), ("left", 110), ("right", 110)):
        mar.append(_el("w:" + side, w=w, type="dxa"))
    tblpr.append(mar)
    tblpr.append(_el("w:tblW", w=5000, type="pct"))
    al = {"right": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}
    last = len(rows) - 1
    for ri, row in enumerate(rows):
        total = ri > 0 and dt.is_total_row(row)
        if ri == 0:
            trpr = tbl.rows[0]._tr.get_or_add_trPr()
            trpr.append(_el("w:tblHeader"))
        for ci in range(ncol):
            cell = tbl.cell(ri, ci)
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            par.alignment = al[align[ci]]
            text = re.sub(r"^\*\*(.*)\*\*$", r"\1", row[ci]) if total else row[ci]
            _runs(par, text, p["on_band"] if ri == 0 else None, 9.5)
            if ri == 0:
                _shade(cell, p["band"])
                for r in par.runs:
                    r.bold = True
                _cell_border(cell, bottom=(0, None))
            elif total:
                for r in par.runs:
                    r.bold = True
                _cell_border(cell, top=(12, p["title"]), bottom=(8, p["title"]))
            else:
                if ri % 2 == 0:
                    _shade(cell, p["surface"])
                _cell_border(cell, bottom=(4 if ri < last else 8, p["line"] if ri < last else p["title"]))
    # keep rows from splitting across pages
    for tr in tbl._tbl.findall(qn("w:tr")):
        tr.get_or_add_trPr().append(_el("w:cantSplit"))
    return tbl


def make_docx(path, markdown, title=None, theme_name=None, subtitle=None):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
    t = theme(theme_name)
    p = paper(t)
    hfont, bfont = p["fonts"]
    doc = Document()
    dt.apply_office_theme(doc.part.package, t)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(1)
        sec.top_margin, sec.bottom_margin = Inches(0.9), Inches(0.85)
    _docx_styles(doc, p)
    if title:
        doc.core_properties.title = title
        band = doc.add_paragraph()
        band.paragraph_format.space_after = Pt(10)
        _font(band.add_run("█" * 7), "Segoe UI", 5, p["accent"])      # a short accent bar
        doc.add_paragraph(title, style="Title")
        sub = doc.add_paragraph(subtitle or "", style="Subtitle")
        _par_border(sub, bottom=(6, p["line"], 10))
    warnings, charts, tables = [], 0, 0
    blocks = dt.parse_blocks(markdown)
    for b in blocks:
        k = b["kind"]
        if k == "heading":
            doc.add_heading(b["text"], level=b["level"])
        elif k == "para":
            _runs(doc.add_paragraph(), b["text"], link=p["accent"])
        elif k in ("bullets", "numbered"):
            for level, text in b["items"]:
                style = "List Number" if k == "numbered" else ("List Bullet 2" if level else "List Bullet")
                par = doc.add_paragraph(style=style)
                _runs(par, text)
        elif k == "table":
            _docx_table(doc, b["rows"], b["align"], p)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            tables += 1
        elif k == "quote":
            par = doc.add_paragraph()
            pf = par.paragraph_format
            pf.left_indent, pf.right_indent = Pt(4), Pt(4)
            pf.space_before, pf.space_after = Pt(6), Pt(10)
            _shade(par, dt.tint(p, 0.9))
            _par_border(par, left=(24, p["accent"], 10), top=(2, dt.tint(p, 0.9), 6), bottom=(2, dt.tint(p, 0.9), 6),
                        right=(2, dt.tint(p, 0.9), 6))
            _runs(par, b["text"], p["text"])
        elif k == "code":
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Pt(4)
            par.paragraph_format.line_spacing = 1.0
            _shade(par, p["surface"])
            r = par.add_run(b["text"])
            _font(r, "Consolas", 9, p["text"])
        elif k == "chart":
            try:
                png = dt.chart_image(b["spec"], t)
            except Exception as e:
                warnings.append(f"chart skipped ({e})")
                continue
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.add_run().add_picture(io.BytesIO(png), width=Inches(6.2))
            charts += 1
        elif k == "image":
            if os.path.isfile(b["path"]):
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.add_run().add_picture(b["path"], width=Inches(6.0))
                if b["alt"]:
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _font(cap.add_run(b["alt"]), size=9, color=p["muted"], italic=True)
            else:
                warnings.append(f"image not found: {b['path']}")
        elif k == "pagebreak":
            doc.add_page_break()
    # header (title, from page 2) and footer (page number on every page)
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    if title:
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _font(hp.add_run(title), size=8.5, color=p["muted"])
    for footer in (sec.footer, sec.first_page_footer):
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _font(fp.add_run("Page "), size=8.5, color=p["muted"])
        _field(fp, "PAGE", 8.5, p["muted"])
    doc.save(path)
    info = (f"{sum(1 for x in doc.paragraphs if x.text.strip())} paragraphs, {tables} tables, {charts} charts, "
            f"theme {dt.theme_name(theme_name)}")
    return info + ("; " + "; ".join(warnings) if warnings else "")


# ================================================================== Excel
def _num(v):
    if isinstance(v, str):
        t = v.strip().replace(",", "")
        if re.fullmatch(r"-?\d+", t):
            return int(t) if len(t.lstrip("-")) < 16 else v
        if re.fullmatch(r"-?\d*\.\d+(e-?\d+)?", t, re.I):
            return float(t)
        if re.fullmatch(r"-?\d+(\.\d+)?%", t):
            return float(t[:-1]) / 100
        m = re.fullmatch(r"(-?)[$€£](-?\d+(\.\d+)?)", t)
        if m:
            return float(m.group(1) + m.group(2))
        m = re.fullmatch(r"\((?:[$€£])?(\d+(\.\d+)?)\)", t)
        if m:
            return -float(m.group(1))
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", t):
            try:
                return datetime.date.fromisoformat(t)
            except ValueError:
                return v
    return v


def _infer_format(header, raw):
    vals = [str(v).strip() for v in raw if v not in (None, "")]
    if not vals:
        return None
    sym = next((s for s in "$€£" if all(v.lstrip("-(").startswith(s) for v in vals)), None)
    if sym:
        dec = any("." in v for v in vals)
        base = f"{sym}#,##0" + (".00" if dec else "")
        return f"{base};[Red]-{base}"
    if all(v.endswith("%") for v in vals):
        return "0.0%" if any("." in v for v in vals) else "0%"
    if all(re.fullmatch(r"\d{4}-\d{2}-\d{2}", v) for v in vals):
        return "yyyy-mm-dd"
    nums = [_num(v) for v in vals]
    if all(isinstance(n, (int, float)) and not isinstance(n, bool) for n in nums):
        if all(isinstance(n, int) for n in nums):
            if "year" in str(header).lower() or all(1800 <= n <= 2200 for n in nums):
                return "0"
            return "#,##0" if any(abs(n) >= 1000 for n in nums) else "0"
        return "#,##0.00" if any(abs(n) >= 1000 for n in nums) else "0.00"
    return None


def _xl_chart(ws, ch, t, header_row, first, last, ncols):
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference, ScatterChart, Series, AreaChart
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.layout import Layout, ManualLayout  # noqa: F401
    from openpyxl.chart.series import DataPoint
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.chart.text import RichText
    from openpyxl.drawing.line import LineProperties
    from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, Font as DFont
    p = paper(t)
    kind = str(ch.get("type") or "column").lower()
    cat_col = int(ch.get("category_col") or 1)
    val_cols = [int(x) for x in (ch.get("value_cols") or [2])]
    cls = {"bar": BarChart, "column": BarChart, "line": LineChart, "pie": PieChart, "donut": PieChart,
           "scatter": ScatterChart, "area": AreaChart}.get(kind, BarChart)
    chart = cls()
    if cls is BarChart:
        chart.type = "bar" if kind == "bar" else "col"
        chart.gapWidth = 60
        chart.overlap = -10 if len(val_cols) > 1 else 0
    if kind == "donut":
        from openpyxl.chart import DoughnutChart
        chart = DoughnutChart()
        chart.holeSize = 55
    chart.title = ch.get("title") or None
    chart.style = None
    chart.roundedCorners = False
    chart.height, chart.width = float(ch.get("height") or 8), float(ch.get("width") or 17)
    cats = Reference(ws, min_col=cat_col, min_row=first, max_row=last)
    palette = p["palette"]
    if kind == "scatter":
        for vc in val_cols:
            chart.series.append(Series(Reference(ws, min_col=vc, min_row=header_row, max_row=last), cats,
                                       title_from_data=True))
    else:
        for vc in val_cols:
            chart.add_data(Reference(ws, min_col=vc, min_row=header_row, max_row=last), titles_from_data=True)
        chart.set_categories(cats)
    for k, s in enumerate(chart.series):
        col = palette[k % len(palette)]
        if kind in ("line", "scatter"):
            s.graphicalProperties.line.solidFill = col
            s.graphicalProperties.line.width = 28575
            s.smooth = False
            s.marker.symbol = "circle"
            s.marker.size = 6
            s.marker.graphicalProperties = GraphicalProperties(solidFill=col)
            s.marker.graphicalProperties.line.solidFill = col
        elif kind in ("pie", "donut"):
            for j in range(last - first + 1):
                pt = DataPoint(idx=j)
                pt.graphicalProperties.solidFill = palette[j % len(palette)]
                pt.graphicalProperties.line.solidFill = "FFFFFF"
                s.dPt.append(pt)
        else:
            s.graphicalProperties.solidFill = col
            s.graphicalProperties.line.noFill = True
    muted = CharacterProperties(sz=900, solidFill=p["muted"], latin=DFont(typeface=p["fonts"][1]))
    txt = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=muted), endParaRPr=muted)])
    if kind in ("pie", "donut"):
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        chart.dataLabels.showVal = chart.dataLabels.showCatName = chart.dataLabels.showSerName = False
        chart.dataLabels.showLeaderLines = False
        chart.legend.position = "r"
    else:
        for ax in (chart.x_axis, chart.y_axis):
            ax.delete = False                              # openpyxl 3.1 hides axes unless told otherwise
            ax.txPr = txt
            ax.spPr = GraphicalProperties(ln=LineProperties(solidFill=p["line"]))
        chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=p["line"]))
        chart.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
        if ch.get("x_title"):
            chart.x_axis.title = ch["x_title"]
        if ch.get("y_title"):
            chart.y_axis.title = ch["y_title"]
        if len(val_cols) == 1:
            chart.legend = None
        else:
            chart.legend.position = "t"
    if chart.legend is not None:
        chart.legend.txPr = txt
    if chart.title is not None:
        from openpyxl.chart.title import Title  # noqa: F401
        chart.title.overlay = False
        head = CharacterProperties(sz=1200, b=True, solidFill=p["title"], latin=DFont(typeface=p["fonts"][0]))
        chart.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=head)
        for r in chart.title.tx.rich.p[0].r or []:
            r.rPr = head
    return chart


def make_xlsx(path, content, theme_name=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.writer.theme import theme_xml
    data = content
    if isinstance(content, str):
        try:
            data = json.loads(content)
        except ValueError:
            data = {"sheets": [{"name": "Sheet1", "rows": list(csv.reader(io.StringIO(content)))}]}
    if isinstance(data, list):
        data = {"sheets": [{"name": "Sheet1", "rows": data}]}
    tname = data.get("theme") or theme_name
    if tname and str(tname).lower() not in THEMES:
        raise DocError(f"unknown theme '{tname}'; use one of: {', '.join(THEMES)}")
    t = theme(tname)
    p = paper(t)
    hfont, bfont = p["fonts"]
    sheets = data.get("sheets") or []
    if not sheets:
        raise DocError('xlsx content needs {"sheets": [{"name": ..., "rows": [[header...], [row...]]}]} or CSV text')
    wb = Workbook()
    wb.remove(wb.active)
    wb.loaded_theme = dt.office_theme_xml(theme_xml, t)
    total, charts, table_names = 0, 0, set()
    line = Side(style="thin", color=p["line"])
    strong = Side(style="medium", color=p["title"])
    for sh in sheets:
        ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "-", str(sh.get("name") or f"Sheet{len(wb.sheetnames) + 1}"))[:31])
        ws.sheet_properties.tabColor = p["accent"]
        ws.sheet_view.showGridLines = False
        rows = [list(r) if isinstance(r, (list, tuple)) else [r] for r in (sh.get("rows") or [])]
        if sh.get("rows") and isinstance(sh["rows"][0], dict):
            keys = list(sh["rows"][0].keys())
            rows = [keys] + [[r.get(k) for k in keys] for r in sh["rows"]]
        header = bool(rows) and sh.get("header", True)
        ncols = max((len(r) for r in rows), default=0)
        rows = [r + [None] * (ncols - len(r)) for r in rows]
        top = 1
        if sh.get("title"):
            c = ws.cell(1, 1, str(sh["title"]))
            c.font = Font(name=hfont, size=16, bold=True, color=p["title"])
            ws.row_dimensions[1].height = 26
            top = 2
            if sh.get("subtitle"):
                c = ws.cell(2, 1, str(sh["subtitle"]))
                c.font = Font(name=bfont, size=10, color=p["muted"])
                top = 3
            top += 1
        if header:
            hdr, seen = [], set()
            for ci, h in enumerate(rows[0], 1):
                h = str(h if h not in (None, "") else f"Column {ci}")
                while h in seen:
                    h += " "
                seen.add(h)
                hdr.append(h)
            rows[0] = hdr
        body = rows[1:] if header else rows
        has_total = header and len(body) > 1 and dt.is_total_row([str(x or "") for x in body[-1]])
        for ri, r in enumerate(rows):
            for ci, v in enumerate(r, 1):
                cell = ws.cell(top + ri, ci, v if (header and ri == 0) else _num(v))
                cell.font = Font(name=bfont, size=10, color=p["text"])
        total += max(0, len(body) - (1 if has_total else 0))
        first, last = top + (1 if header else 0), top + len(rows) - 1
        data_last = last - (1 if has_total else 0)
        # number formats: explicit per column, else inferred from the source text ($ / % / dates / big ints)
        fmts = dict(sh.get("formats") or {})
        numeric_cols = set()
        for ci in range(1, ncols + 1):
            letter = get_column_letter(ci)
            raw = [r[ci - 1] for r in body]
            if letter not in fmts:
                f = _infer_format(rows[0][ci - 1] if header else "", raw)
                if f:
                    fmts[letter] = f
            if all(isinstance(_num(v), (int, float, datetime.date)) for v in raw if v not in (None, "")) and \
                    any(v not in (None, "") for v in raw):
                numeric_cols.add(ci)
        for letter, f in fmts.items():
            for ri in range(first, last + 1):
                ws[f"{letter}{ri}"].number_format = f
        if sh.get("totals") and header and body and not has_total:
            last += 1
            has_total = True
            ws.cell(last, 1, "Total").font = Font(name=bfont, size=10)
            for ci in sorted(numeric_cols):
                letter = get_column_letter(ci)
                if ci == 1 or "%" in fmts.get(letter, "") or fmts.get(letter) in ("0", "yyyy-mm-dd") and \
                        "year" in str(rows[0][ci - 1]).lower():
                    continue
                c = ws.cell(last, ci, f"=SUM({letter}{first}:{letter}{data_last})")
                c.number_format = fmts.get(letter, "General")
        for ci in range(1, ncols + 1):
            for ri in range(first, last + 1):
                c = ws.cell(ri, ci)
                c.alignment = Alignment(horizontal="right" if ci in numeric_cols else "left", vertical="center")
        if header:
            for ci in range(1, ncols + 1):
                c = ws.cell(top, ci)
                c.font = Font(name=hfont, size=10, bold=True, color=p["on_band"])
                c.fill = PatternFill("solid", fgColor=p["band"])
                c.alignment = Alignment(vertical="center", horizontal="right" if ci in numeric_cols else "left",
                                        wrap_text=True, indent=0)
            ws.row_dimensions[top].height = 24
            ws.freeze_panes = ws.cell(top + 1, 1).coordinate
            ws.print_title_rows = f"{top}:{top}"
            if len(body) >= 1 and data_last >= first:
                name = re.sub(r"\W", "_", ws.title) or "Table"
                name = ("T_" + name if not name[0].isalpha() else name)[:200]
                while name in table_names:
                    name += "_"
                table_names.add(name)
                tab = Table(displayName=name, ref=f"A{top}:{get_column_letter(ncols)}{data_last}")
                tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
                ws.add_table(tab)
            for ri in range(first, data_last + 1):
                for ci in range(1, ncols + 1):
                    c = ws.cell(ri, ci)
                    c.border = Border(bottom=line)
                    if (ri - first) % 2 == 1:
                        c.fill = PatternFill("solid", fgColor=p["surface"])
            if has_total:
                for ci in range(1, ncols + 1):
                    c = ws.cell(last, ci)
                    c.font = Font(name=bfont, size=10, bold=True, color=p["title"])
                    c.border = Border(top=strong, bottom=strong)
        for ci in range(1, ncols + 1):
            letter = get_column_letter(ci)
            width = 8
            for ri in range(top, last + 1):
                v = ws.cell(ri, ci).value
                if v is None or (sh.get("title") and ri < top):
                    continue
                s = f"{v:,.2f}" if isinstance(v, float) else (f"{v:,}" if isinstance(v, int) else str(v))
                width = max(width, len(s) + (2 if ri == top else 1))
            ws.column_dimensions[letter].width = min(60, width + 3)
        # native charts, below the table so they print with it and never cover data
        anchor_row = last + 3
        for ch in sh.get("charts") or []:
            if not header or data_last < first:
                break
            chart = _xl_chart(ws, ch, t, top, first, data_last, ncols)
            ws.add_chart(chart, ch.get("anchor") or f"A{anchor_row}")
            anchor_row += int(float(ch.get("height") or 8) * 2) + 2
            charts += 1
        # print setup: fit to page width, landscape for wide sheets, page numbers
        ws.page_setup.orientation = "landscape" if ncols > 6 or sh.get("charts") else "portrait"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.oddFooter.left.text = "&A"
        ws.oddFooter.right.text = "Page &P of &N"
        ws.oddFooter.left.size = ws.oddFooter.right.size = 8
    wb.save(path)
    return f"{len(sheets)} sheet(s), {total} data rows, {charts} chart(s), theme {dt.theme_name(tname)}"


# ================================================================== PowerPoint
def _pp_color(hexs):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hexs)


def _box(slide, x, y, w, h, fill=None, shape="rect", line=None):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Pt
    kind = {"rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE, "oval": MSO_SHAPE.OVAL}[shape]
    shp = slide.shapes.add_shape(kind, int(x), int(y), int(w), int(h))
    if shape == "round":
        shp.adjustments[0] = 0.06
    if line:
        shp.line.color.rgb = _pp_color(line[0])
        shp.line.width = Pt(line[1])
    else:
        shp.line.fill.background()
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = _pp_color(fill)
    else:
        shp.fill.background()
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def _fit(lines, w_emu, h_emu, size, minimum=12, spacing=1.35):
    """Shrink a font size until the text (list of strings) fits the box, by a simple character-width estimate."""
    import math
    w_in, h_in = w_emu / 914400, h_emu / 914400
    while size > minimum:
        cpl = max(8, (w_in * 72) / (size * 0.52))
        n = sum(max(1, math.ceil(len(str(s)) / cpl)) for s in lines)
        if n * size * spacing / 72 + 0.1 * len(lines) * size / 72 <= h_in:
            break
        size -= 1
    return size


def _md_plain(s):
    return re.sub(r"\*\*([^*]+)\*\*|__([^_]+)__|`([^`]+)`", lambda m: m.group(1) or m.group(2) or m.group(3), str(s))


def _text(slide, x, y, w, h, text, size, color, font, bold=False, align=None, anchor=None, italic=False, fit=False):
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Pt
    tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.vertical_anchor = {"middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}.get(anchor, MSO_ANCHOR.TOP)
    lines = text if isinstance(text, list) else [text]
    if fit:
        size = _fit([_md_plain(x) for x in lines], w, h, size)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(align, PP_ALIGN.LEFT)
        r = p.add_run()
        r.text = _md_plain(line)
        r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), bold, italic, font
        r.font.color.rgb = _pp_color(color)
    return tb


def _rich_line(p, text, size, t, color=None, label_color=None):
    """'Label: rest' -> bold label; **bold** spans honoured."""
    from pptx.util import Pt
    text = str(text)
    m = re.match(r"^\s*(\*\*)?([^:*]{2,40}?)(\*\*)?:\s+(.+)$", text)
    parts = [(m.group(2) + ": ", True), (m.group(4), False)] if m and (m.group(1) or len(m.group(2)) <= 28) else \
        [(x, i % 2 == 1) for i, x in enumerate(re.split(r"\*\*([^*]+)\*\*", text))]
    for s, bold in parts:
        if not s:
            continue
        r = p.add_run()
        r.text = _md_plain(s)
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, t["fonts"][1]
        r.font.color.rgb = _pp_color((label_color or t["title"]) if bold else (color or t["text"]))


def _bullets(slide, x, y, w, h, items, t, size=22, fit=True):
    from pptx.util import Pt
    items = list(items or [])
    flat = [(int(b.get("level") or 0), b.get("text", "")) if isinstance(b, dict) else (0, b) for b in items]
    if fit:
        size = _fit([_md_plain(s) + "    " for _, s in flat], w - 400000, h, size, minimum=14, spacing=1.6)
    tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    for i, (level, b) in enumerate(flat):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        sz = size - 4 * min(level, 2)
        p.space_after = Pt(round(sz * 0.6))
        p.level = min(level, 3)
        dot = p.add_run()
        dot.text = ("●   " if level == 0 else "–   ")
        dot.font.color.rgb = _pp_color(t["accent"])
        dot.font.size = Pt(round(sz * (0.55 if level == 0 else 1)))
        _rich_line(p, b, sz, t)
    return tb


def _md_slides(md):
    slides, cur = [], None
    levels = [len(m.group(1)) for m in re.finditer(r'^\s*(#{1,3})\s', md, re.M)]
    slide_level = min(levels) if levels else 1
    for ln in md.replace("\r\n", "\n").split("\n"):
        s = ln.strip()
        numbered = re.match(r"^Slide\s+\d+\s*:\s*(.*)$", s.strip('*').strip(), re.I)
        if numbered:
            cur = {"layout": "bullets", "title": numbered.group(1).strip(), "bullets": [], "notes": ""}
            slides.append(cur)
        elif cur is not None and s.lower().startswith("title:"):
            cur["title"] = s[6:].strip()
        elif cur is not None and s.lower() == "bullets:":
            continue
        elif re.match(r"^#{" + str(slide_level) + r"}\s", s):
            cur = {"layout": "bullets", "title": s.lstrip("#").strip(), "bullets": [], "notes": ""}
            slides.append(cur)
        elif cur is not None and s.startswith("|"):
            if not dt._TABLE_SEP.match(s):
                cur.setdefault("rows", []).append([c.strip() for c in s.strip("|").split("|")])
                cur["layout"] = "table"
        elif cur is not None and re.match(r"^[-*+]\s|^\d+[.)]\s", s):
            cur["bullets"].append(re.sub(r"^([-*+]|\d+[.)])\s", "", s))
        elif cur is not None and s.lower().startswith("notes:"):
            cur["notes"] = s[6:].strip()
        elif cur is not None and s:
            cur["bullets"].append(re.sub(r'^#{1,3}\s+', '', s))
    return slides


def _card_items(sd):
    """bullets like 'Label: text' (2-4 of them) read better as cards."""
    cards = sd.get("cards")
    if cards:
        return [c if isinstance(c, dict) else {"title": str(c)} for c in cards][:4]
    bl = sd.get("bullets") or []
    if 2 <= len(bl) <= 4 and all(isinstance(b, str) and re.match(r"^\**[^:]{2,32}\**:\s+\S", b) for b in bl):
        out = []
        for b in bl:
            k, v = b.split(":", 1)
            out.append({"title": k.strip("* "), "text": v.strip()})
        return out
    return None


def _pp_table(s, x, y, w, h, rows, t):
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.util import Emu, Pt
    rows = [[str(c) for c in r] for r in rows if r]
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    align = dt.table_align(rows)
    n = len(rows)
    size = 20 if n <= 5 else 18 if n <= 7 else 16 if n <= 9 else 13 if n <= 13 else 11
    row_h = Emu(int(min(h / n, Pt(size * 2.4))))
    shp = s.shapes.add_table(n, ncol, int(x), int(y), int(w), int(row_h * n))
    tbl = shp.table
    tbl.first_row = True
    tbl.horz_banding = False
    tblPr = shp._element.graphic.graphicData.tbl.tblPr
    style = tblPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId")
    if style is not None:
        style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # "No Style, No Grid": every cell is painted here
    first_w = max(len(r[0]) for r in rows)
    other_w = max([len(r[c]) for r in rows for c in range(1, ncol)] or [1])
    share = min(0.5, max(1.0 / ncol, first_w / (first_w + other_w * (ncol - 1) + 1)))
    tbl.columns[0].width = int(w * share)
    for c in range(1, ncol):
        tbl.columns[c].width = int(w * (1 - share) / (ncol - 1))
    surface = t["surface"]
    for ri, r in enumerate(rows):
        tbl.rows[ri].height = row_h
        total = ri > 0 and dt.is_total_row(r)
        for ci, val in enumerate(r):
            cell = tbl.cell(ri, ci)
            cell.margin_left = cell.margin_right = Emu(91440)
            cell.margin_top = cell.margin_bottom = Emu(27432)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = _pp_color(t["band"] if ri == 0 else (surface if ri % 2 == 0 else t["bg"]))
            if total:
                cell.fill.fore_color.rgb = _pp_color(mix(t["accent"], t["bg"], 0.82))
            tf = cell.text_frame
            p = tf.paragraphs[0]
            p.alignment = {"right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}.get(align[ci], PP_ALIGN.LEFT)
            run = p.add_run()
            run.text = _md_plain(val)
            run.font.size = Pt(size)
            run.font.name = t["fonts"][0] if ri == 0 else t["fonts"][1]
            run.font.bold = ri == 0 or total
            run.font.color.rgb = _pp_color(t["on_band"] if ri == 0 else (t["title"] if total else t["text"]))
    return shp


def _pp_chart(s, x, y, w, h, sd, t):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
    from pptx.util import Pt
    spec = dt.normalize_chart({k: sd.get(k) for k in ("type", "labels", "series", "values", "name") if sd.get(k) is not None})
    if not spec["labels"] and sd.get("rows"):
        spec = dt.chart_from_table(spec, [[str(c) for c in r] for r in sd["rows"]])
    if not spec["labels"] or not spec["series"]:
        raise DocError('chart slide needs "labels" and "series": [{"name": ..., "values": [...]}] (or "rows")')
    kind = spec["type"]
    cd = CategoryChartData()
    cd.categories = spec["labels"]
    series = spec["series"][:1] if kind in ("pie", "donut") else spec["series"]
    for se in series:
        cd.add_series(se["name"] or " ", [v if v is not None else 0 for v in se["values"]])
    xl = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS,
          "pie": XL_CHART_TYPE.PIE, "donut": XL_CHART_TYPE.DOUGHNUT, "area": XL_CHART_TYPE.AREA,
          "stacked": XL_CHART_TYPE.COLUMN_STACKED}.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED)
    gf = s.shapes.add_chart(xl, int(x), int(y), int(w), int(h), cd)
    chart = gf.chart
    chart.font.size = Pt(13)
    chart.font.name = t["fonts"][1]
    chart.font.color.rgb = _pp_color(t["muted"])
    pal = t["palette"]
    if kind in ("pie", "donut"):
        plot = chart.plots[0]
        for j, pt in enumerate(plot.series[0].points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = _pp_color(pal[j % len(pal)])
            pt.format.line.color.rgb = _pp_color(t["bg"])
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format, dl.number_format_is_linked = "0%", False
        dl.show_percentage, dl.show_value = True, False
        dl.font.size, dl.font.bold = Pt(14), True
        dl.font.color.rgb = _pp_color("FFFFFF")
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.RIGHT
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(14)
        chart.legend.font.color.rgb = _pp_color(t["text"])
        return gf
    plot = chart.plots[0]
    if hasattr(plot, "gap_width"):
        plot.gap_width = 60
        if len(series) > 1 and kind != "stacked":
            plot.overlap = -10
    for k, se in enumerate(plot.series):
        col = _pp_color(pal[k % len(pal)])
        if kind == "line":
            se.format.line.color.rgb = col
            se.format.line.width = Pt(3)
            se.smooth = False
            se.marker.format.fill.solid()
            se.marker.format.fill.fore_color.rgb = col
            se.marker.format.line.color.rgb = col
        else:
            se.format.fill.solid()
            se.format.fill.fore_color.rgb = col
    va, ca = chart.value_axis, chart.category_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = _pp_color(t["line"])
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(12)
    ca.format.line.color.rgb = _pp_color(t["line"])
    ca.tick_labels.font.size = Pt(13)
    ca.tick_labels.font.color.rgb = _pp_color(t["text"])
    if len(series) == 1 and kind in ("column", "bar") and len(spec["labels"]) <= 12:
        plot.has_data_labels = True
        plot.data_labels.font.size = Pt(13)
        plot.data_labels.font.color.rgb = _pp_color(t["text"])
        plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(13)
        chart.legend.font.color.rgb = _pp_color(t["text"])
    return gf


def make_pptx(path, content, title=None, theme_name=None):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.util import Emu, Inches, Pt
    data = content
    markdown_outline = False
    if isinstance(content, str):
        try:
            data = json.loads(content)
        except ValueError:
            markdown_outline = True
            data = {"slides": _md_slides(content)}
    if isinstance(data, list):
        data = {"slides": data}
    slides = [s for s in (data.get("slides") or []) if isinstance(s, dict)]
    if not slides:
        raise DocError('pptx content needs {"slides": [{"layout": "...", "title": ..., ...}]} or Markdown with # per slide '
                       "(read_guide('powerpoint') has the full format)")
    tname = data.get("theme") or theme_name
    if tname and str(tname).lower() not in THEMES:
        raise DocError(f"unknown theme '{tname}'; use one of: {', '.join(THEMES)}")
    t = theme(tname)
    hfont, bfont = t["fonts"]
    title = title or data.get("title")
    if title and not markdown_outline and not any(s.get("layout") == "title" for s in slides):
        slides.insert(0, {"layout": "title", "title": title, "subtitle": data.get("subtitle", "")})
    deck_title = title or next((s.get("title") for s in slides if s.get("layout") == "title"), "")
    prs = Presentation()
    dt.apply_office_theme(prs.part.package, t)
    prs.core_properties.title = deck_title or ""
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    M = Inches(0.75)
    blank = prs.slide_layouts[6]
    total = len(slides)
    section_no = 0
    on_dark = t["dark"]
    for idx, sd in enumerate(slides, 1):
        lay = str(sd.get("layout") or ("bullets" if sd.get("bullets") else "section")).lower()
        if lay == "bullets" and _card_items(sd):
            lay = "cards"
        if lay == "bullets" and sd.get("rows"):
            lay = "table"
        s = prs.slides.add_slide(blank)
        if lay in ("title", "section", "closing"):
            _box(s, 0, 0, W, H, t["band"])
            ring = mix(t["band"], "FFFFFF", 0.06)
            if lay == "title":
                _box(s, W - Inches(5.2), -Inches(2.2), Inches(8.4), Inches(8.4), ring, "oval")
                _box(s, W - Inches(3.1), H - Inches(2.6), Inches(4.2), Inches(4.2), mix(t["band"], t["accent"], 0.35), "oval")
                _text(s, M, Inches(1.4), W * 0.62, Inches(2.75), sd.get("title", ""), 50, t["on_band"], hfont, True,
                      anchor="bottom", fit=True)
                _box(s, M, Inches(4.38), Inches(0.9), Emu(64008), t["accent"])
                if sd.get("subtitle"):
                    _text(s, M, Inches(4.62), W * 0.6, Inches(1.1), sd["subtitle"], 22, mix(t["on_band"], t["band"], 0.25),
                          bfont, fit=True)
                foot = sd.get("footer") or data.get("footer") or data.get("date")
                if foot:
                    _text(s, M, H - Inches(0.95), W * 0.6, Inches(0.4), foot, 13, t["accent2"], bfont)
            elif lay == "section":
                section_no += 1
                _box(s, -Inches(2.4), Inches(1.0), Inches(5.5), Inches(5.5), ring, "oval")
                _text(s, M + Inches(0.2), Inches(2.0), Inches(3), Inches(1.4), f"{section_no:02d}", 72, t["accent"], hfont, True)
                _text(s, M + Inches(3.0), Inches(2.35), W - M * 2 - Inches(3.2), Inches(1.4), sd.get("title", ""), 40,
                      t["on_band"], hfont, True, anchor="middle", fit=True)
                if sd.get("subtitle"):
                    _text(s, M + Inches(3.0), Inches(3.8), W - M * 2 - Inches(3.2), Inches(1.0), sd["subtitle"], 20,
                          mix(t["on_band"], t["band"], 0.3), bfont)
            else:
                _box(s, W / 2 - Inches(4.5), H / 2 - Inches(4.5), Inches(9), Inches(9), ring, "oval")
                _text(s, M, Inches(2.5), W - 2 * M, Inches(1.3), sd.get("title", "Thank you"), 48, t["on_band"], hfont,
                      True, align="center", anchor="bottom", fit=True)
                _box(s, W / 2 - Inches(0.45), Inches(4.0), Inches(0.9), Emu(64008), t["accent"])
                if sd.get("subtitle"):
                    _text(s, M, Inches(4.25), W - 2 * M, Inches(1.0), sd["subtitle"], 22,
                          mix(t["on_band"], t["band"], 0.25), bfont, align="center")
                if sd.get("bullets"):
                    _text(s, M, Inches(5.1), W - 2 * M, Inches(1.4), [str(b) for b in sd["bullets"]][:4], 16,
                          t["accent2"], bfont, align="center")
        else:
            _box(s, 0, 0, W, H, t["bg"])
            _text(s, M, Inches(0.45), W - 2 * M, Inches(0.95), sd.get("title", ""), 32, t["title"], hfont, True,
                  anchor="bottom", fit=True)
            _box(s, M, Inches(1.5), Inches(0.7), Emu(54864), t["accent"])
            top, bottom = Inches(1.85), H - Inches(0.85)
            body_h = bottom - top
            bw = W - 2 * M
            card_bg = t["surface"]
            if lay == "two_column":
                gap = Inches(0.45)
                colw = (bw - gap) / 2
                for j, side in enumerate(("left", "right")):
                    x = M + j * (colw + gap)
                    col = sd.get(side) or {}
                    if isinstance(col, list):
                        col = {"bullets": col}
                    import math
                    nb = col.get("bullets") or []
                    est = sum(max(1, math.ceil(len(_md_plain(b.get("text", "") if isinstance(b, dict) else b)) /
                                                max(10, (colw - Inches(0.7)) / 914400 * 72 / (20 * 0.52)))) for b in nb)
                    card_h = min(body_h, Inches(0.6 + (0.75 if col.get("heading") else 0) + est * 20 * 1.65 / 72 + 0.3))
                    card_h = max(card_h, min(body_h, Inches(3.2)))
                    _box(s, x, top, colw, card_h, card_bg, "round")
                    _box(s, x, top, colw, Emu(64008), t["palette"][j % 6])
                    pad = Inches(0.35)
                    y0 = top + Inches(0.3)
                    if col.get("heading"):
                        _text(s, x + pad, y0, colw - 2 * pad, Inches(0.6), col["heading"], 22, t["title"], hfont, True)
                        y0 += Inches(0.75)
                    _bullets(s, x + pad, y0, colw - 2 * pad, top + card_h - y0 - Inches(0.2), col.get("bullets"), t, 20)
            elif lay == "cards":
                cards = _card_items(sd) or []
                n = max(1, len(cards))
                gap = Inches(0.35)
                cw = (bw - gap * (n - 1)) / n
                ch = min(body_h, Inches(3.9))
                y = top + (body_h - ch) / 2 - Inches(0.1)
                tsize = min([_fit([c.get("title", "")], cw - Inches(0.7), Inches(0.9), 24, 14) for c in cards] or [24])
                bsize = min([_fit([c.get("text", "")], cw - Inches(0.7), ch - Inches(2.1), 18, 12) for c in cards] or [18])
                for j, c in enumerate(cards):
                    x = M + j * (cw + gap)
                    col = t["palette"][j % 6]
                    _box(s, x, y, cw, ch, card_bg, "round")
                    _box(s, x + Inches(0.35), y + Inches(0.4), Inches(0.55), Inches(0.55), col, "oval")
                    _text(s, x + Inches(0.35), y + Inches(0.4), Inches(0.55), Inches(0.55), str(c.get("icon") or j + 1), 16,
                          "FFFFFF", hfont, True, align="center", anchor="middle")
                    _text(s, x + Inches(0.35), y + Inches(1.15), cw - Inches(0.7), Inches(0.9), c.get("title", ""), tsize,
                          t["title"], hfont, True)
                    _text(s, x + Inches(0.35), y + Inches(2.0), cw - Inches(0.7), ch - Inches(2.2), c.get("text", ""),
                          bsize, t["text"], bfont)
            elif lay == "stats":
                stats = sd.get("stats") or []
                n = max(1, len(stats))
                gap = Inches(0.35)
                cw = (bw - gap * (n - 1)) / n
                ch = Inches(2.9)
                y = top + max(Inches(0.2), (body_h - ch - (Inches(1.0) if sd.get("caption") else 0)) / 2 - Inches(0.2))
                vsize = min([_fit([str(st.get("value", ""))], cw - Inches(0.6), Inches(1.3), 60, 28) for st in stats] or [60])
                for j, st in enumerate(stats):
                    x = M + j * (cw + gap)
                    col = t["palette"][j % 6]
                    _box(s, x, y, cw, ch, card_bg, "round")
                    _box(s, x, y, cw, Emu(73152), col)
                    _text(s, x + Inches(0.3), y + Inches(0.45), cw - Inches(0.6), Inches(1.3), str(st.get("value", "")),
                          vsize, col if not on_dark else t["title"], hfont, True, anchor="middle")
                    _text(s, x + Inches(0.3), y + Inches(1.85), cw - Inches(0.6), Inches(0.9), str(st.get("label", "")),
                          17, t["text"], bfont, fit=True)
                if sd.get("caption"):
                    _text(s, M, y + ch + Inches(0.45), bw, Inches(0.8), sd["caption"], 15, t["muted"], bfont)
            elif lay == "quote":
                _text(s, M, top - Inches(0.2), Inches(1.5), Inches(1.6), "“", 140, t["accent"], "Georgia", True)
                _text(s, M + Inches(1.3), top + Inches(0.3), bw - Inches(2.0), Inches(3.0), sd.get("quote", ""), 32,
                      t["title"], "Georgia" if hfont != "Georgia" else hfont, italic=True, anchor="middle", fit=True)
                if sd.get("author"):
                    _text(s, M + Inches(1.3), top + Inches(3.5), bw - Inches(2.0), Inches(0.6), f"— {sd['author']}", 18,
                          t["accent"], bfont, True)
            elif lay == "table" and sd.get("rows"):
                side_text = sd.get("bullets") or []
                tw = bw * (0.64 if side_text else 1.0)
                _pp_table(s, M, top, tw, body_h, sd["rows"], t)
                if side_text:
                    _bullets(s, M + tw + Inches(0.5), top, bw - tw - Inches(0.5), body_h, side_text, t, 18)
                if sd.get("caption"):
                    _text(s, M, bottom - Inches(0.35), bw, Inches(0.35), sd["caption"], 12, t["muted"], bfont)
            elif lay == "chart":
                side_text = sd.get("bullets") or []
                cw = bw * (0.62 if side_text else 1.0)
                _pp_chart(s, M - Inches(0.1), top, cw, body_h - (Inches(0.4) if sd.get("caption") else 0), sd, t)
                if side_text:
                    _bullets(s, M + cw + Inches(0.4), top + Inches(0.3), bw - cw - Inches(0.4), body_h, side_text, t, 20)
                if sd.get("caption"):
                    _text(s, M, bottom - Inches(0.35), bw, Inches(0.35), sd["caption"], 12, t["muted"], bfont)
            elif lay == "timeline":
                steps = [st if isinstance(st, dict) else {"label": str(st)} for st in (sd.get("steps") or sd.get("bullets") or [])][:6]
                n = max(1, len(steps))
                yline = top + body_h * 0.4
                seg = bw / n
                ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(M), int(yline), int(M + bw), int(yline))
                ln.line.color.rgb = _pp_color(t["line"])
                ln.line.width = Pt(3)
                for j, st in enumerate(steps):
                    cx = M + seg * j + seg / 2
                    col = t["palette"][j % 6]
                    _box(s, cx - Inches(0.22), yline - Inches(0.22), Inches(0.44), Inches(0.44), col, "oval",
                         line=(t["bg"], 4))
                    _text(s, cx - seg / 2 + Inches(0.1), yline - Inches(1.35), seg - Inches(0.2), Inches(0.95),
                          st.get("label", ""), 22, t["title"], hfont, True, align="center",
                          anchor="bottom", fit=True)
                    if st.get("text"):
                        _text(s, cx - seg / 2 + Inches(0.15), yline + Inches(0.5), seg - Inches(0.3), Inches(1.8),
                              st["text"], 18, t["text"], bfont, align="center", fit=True)
            elif lay == "image" and sd.get("image") and os.path.isfile(sd["image"]):
                has_text = bool(sd.get("bullets"))
                img_w = bw * (0.56 if has_text else 1.0)
                pic = s.shapes.add_picture(sd["image"], int(M), int(top), width=int(img_w))
                if pic.height > body_h:
                    ratio = body_h / pic.height
                    pic.height, pic.width = int(pic.height * ratio), int(pic.width * ratio)
                if not has_text:
                    pic.left = int((W - pic.width) / 2)
                if has_text:
                    _bullets(s, M + img_w + Inches(0.45), top, bw - img_w - Inches(0.45), body_h, sd["bullets"], t, 20)
                if sd.get("caption"):
                    _text(s, M, bottom - Inches(0.35), bw, Inches(0.35), sd["caption"], 12, t["muted"], bfont)
            else:
                items = sd.get("bullets") or ([sd["body"]] if sd.get("body") else [])
                flat = all(isinstance(b, str) or not int(b.get("level") or 0) for b in items)
                if 2 <= len(items) <= 5 and flat and not sd.get("font_size"):
                    texts = [b if isinstance(b, str) else b.get("text", "") for b in items]
                    row_h = min(body_h / len(texts), Inches(1.15))
                    size = min([_fit([_md_plain(x)], bw - Inches(0.9), row_h - Inches(0.2), 26, 16) for x in texts])
                    for j, x in enumerate(texts):
                        y = top + Inches(0.1) + row_h * j
                        _box(s, M, y + row_h / 2 - Inches(0.07), Inches(0.14), Inches(0.14), t["palette"][j % 6] if
                             t["palette"][j % 6] != t["bg"] else t["accent"], "oval")
                        tb = _text(s, M + Inches(0.45), y, bw - Inches(0.45), row_h, "", size, t["text"], bfont,
                                   anchor="middle")
                        para = tb.text_frame.paragraphs[0]
                        para.runs[0].text = ""
                        _rich_line(para, x, size, t)
                        if j < len(texts) - 1:
                            ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(M + Inches(0.45)), int(y + row_h),
                                                        int(M + bw), int(y + row_h))
                            ln.line.color.rgb = _pp_color(t["line"])
                            ln.line.width = Pt(1)
                elif items:
                    _bullets(s, M, top + Inches(0.1), bw * (0.9 if len(items) > 2 else 1), body_h,
                             items, t, int(sd.get("font_size") or 26))
            # footer: deck title left, slide number right
            if deck_title:
                _text(s, M, H - Inches(0.55), bw * 0.7, Inches(0.35), deck_title, 11, t["muted"], bfont)
            _text(s, W - M - Inches(1.5), H - Inches(0.55), Inches(1.5), Inches(0.35), f"{idx} / {total}", 11,
                  t["muted"], bfont, align="right")
        if sd.get("notes"):
            s.notes_slide.notes_text_frame.text = str(sd["notes"])
    prs.save(path)
    return f"{len(prs.slides)} slides, theme {dt.theme_name(tname)}, 16:9"


# ================================================================== PDF
def make_pdf(path, markdown, title=None, theme_name=None, subtitle=None):
    import doc_pdf
    try:
        return doc_pdf.make_pdf(path, markdown, title, theme_name, subtitle)
    except doc_pdf.PdfError as e:
        raise DocError(str(e))


def create(path, content, kind=None, title=None, theme_name=None, subtitle=None):
    kind = (kind or os.path.splitext(path)[1].lstrip(".")).lower()
    kind = {"doc": "docx", "word": "docx", "excel": "xlsx", "xls": "xlsx", "powerpoint": "pptx", "ppt": "pptx",
            "slides": "pptx", "deck": "pptx", "presentation": "pptx"}.get(kind, kind)
    if kind not in ("docx", "xlsx", "pptx", "pdf"):
        raise DocError("type must be docx, xlsx, pptx or pdf")
    if theme_name and str(theme_name).lower().strip() not in THEMES:
        raise DocError(f"unknown theme '{theme_name}'; use one of: {', '.join(THEMES)}")
    if not path.lower().endswith("." + kind):
        path = os.path.splitext(path)[0] + "." + kind
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    text = content if isinstance(content, str) else json.dumps(content)
    if kind == "docx":
        info = make_docx(path, text, title, theme_name, subtitle)
    elif kind == "xlsx":
        info = make_xlsx(path, content, theme_name)
    elif kind == "pptx":
        info = make_pptx(path, content, title, theme_name)
    else:
        info = make_pdf(path, text, title, theme_name, subtitle)
    return path, kind, info
